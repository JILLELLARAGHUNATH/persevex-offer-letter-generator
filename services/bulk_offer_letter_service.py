import io
import re
import csv
from datetime import datetime
from pathlib import Path


# ============================================================
# COLUMN HEADER SYNONYMS AND MAPPINGS
# ============================================================

HEADER_MAP = {
    "student_name": [
        "student name", "student_name", "name", "full name", "candidate name",
        "student fullname", "applicant name", "candidatename", "studentname", "fullname"
    ],
    "student_email": [
        "student email", "student_email", "email", "email address", "email id",
        "candidate email", "applicant email", "mail", "mail id", "emailid", "studentemail"
    ],
    "domain": [
        "domain", "internship domain", "internship_domain", "track", "field",
        "stream", "role", "position", "profile"
    ],
    "letter_type": [
        "letter type", "letter_type", "type", "hours type", "letter_template",
        "offer letter type", "template"
    ],
    "duration": [
        "duration", "internship duration", "internship_duration", "months",
        "months of duration", "duration months", "period"
    ],
    "start_date": [
        "start date", "start_date", "from date", "start", "commencement date",
        "joining date", "startdate"
    ],
    "end_date": [
        "end date", "end_date", "to date", "end", "completion date", "enddate"
    ],
    "hours_per_week": [
        "hours per week", "hours_per_week", "hours/week", "weekly hours",
        "hours", "hrs/week", "work hours"
    ],
    "stipend": [
        "stipend", "stipend amount", "compensation", "salary", "stipend (optional)"
    ],
}


def normalize_column_name(header):
    if not header:
        return ""
    h = str(header).strip().lower().replace("_", " ")
    h = re.sub(r"[\s\-\.\:]+", " ", h).strip()

    # Pass 1: Exact match against synonyms
    for canonical, synonyms in HEADER_MAP.items():
        if h in synonyms:
            return canonical

    # Pass 2: Whole word / boundary match (synonym matches as a distinct word in header)
    for canonical, synonyms in HEADER_MAP.items():
        for syn in synonyms:
            if re.search(r"\b" + re.escape(syn) + r"\b", h):
                return canonical

    return h.replace(" ", "_")


def is_valid_email(email_str):
    if not email_str or not isinstance(email_str, str):
        return False
    email_regex = r"^[^@\s]+@[^@\s]+\.[^@\s]+$"
    return bool(re.match(email_regex, str(email_str).strip()))


def clean_stipend(value):
    if value is None:
        return ""
    s = str(value).strip().replace("₹", "").replace(",", "")
    return s.strip()


def parse_and_normalize_date(date_val):
    """
    Parse date value and return ISO formatted 'YYYY-MM-DD' and display 'DD/MM/YYYY'.
    """
    if not date_val:
        return None, None, "Date is missing."
    date_str = str(date_val).strip()

    formats = [
        "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d",
        "%d-%b-%Y", "%d %b %Y", "%d-%B-%Y", "%d %B %Y",
        "%m/%d/%Y", "%m-%d-%Y"
    ]

    for fmt in formats:
        try:
            dt = datetime.strptime(date_str, fmt)
            return dt.strftime("%Y-%m-%d"), dt.strftime("%d/%m/%Y"), None
        except ValueError:
            continue

    if "T" in date_str or " " in date_str:
        try:
            dt = datetime.fromisoformat(date_str.replace("Z", "+00:00")[:10])
            return dt.strftime("%Y-%m-%d"), dt.strftime("%d/%m/%Y"), None
        except Exception:
            pass

    return None, None, f"Invalid date format: '{date_str}'. Expected YYYY-MM-DD or DD/MM/YYYY."


def normalize_duration(duration_val):
    if duration_val is None:
        return "2 months"
    s = str(duration_val).strip().lower()
    match = re.search(r"(\d+)", s)
    if match:
        num = int(match.group(1))
        num = max(2, min(6, num))
        return f"{num} months"
    return "2 months"



def get_months_number(duration_val):
    s = str(duration_val or "").strip().lower()
    match = re.search(r"(\d+)", s)
    if match:
        num = int(match.group(1))
        if 2 <= num <= 6:
            return num
    return 2


def calculate_weeks_from_duration(duration_val):
    months = get_months_number(duration_val)
    return months * 4


def calculate_total_hours(weeks_num, hours_per_week_val):
    try:
        hours = int(hours_per_week_val)
        if hours > 0:
            return weeks_num * hours
    except (ValueError, TypeError):
        pass
    return None


def normalize_letter_type(letter_type_val, hours_per_week_val=None):
    if not letter_type_val:
        if hours_per_week_val and str(hours_per_week_val).strip():
            return "with_hours"
        return "without_hours"
    s = str(letter_type_val).strip().lower().replace(" ", "_")
    if "without" in s or "no_hour" in s or "no" in s:
        return "without_hours"
    if "with" in s or "hour" in s or s == "1":
        return "with_hours"
    return "without_hours"


# ============================================================
# CSV / EXCEL PARSING
# ============================================================

def parse_csv_content(file_bytes):
    text = ""
    for enc in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
        try:
            text = file_bytes.decode(enc)
            break
        except Exception:
            continue
    if not text:
        raise ValueError("Could not decode CSV file. Please ensure it is saved as UTF-8.")

    first_line = text.split("\n")[0] if "\n" in text else text
    delimiter = ","
    if "\t" in first_line:
        delimiter = "\t"
    elif ";" in first_line:
        delimiter = ";"

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        return []

    headers = [normalize_column_name(h) for h in rows[0]]
    parsed_rows = []
    for r in rows[1:]:
        if not any(str(c).strip() for c in r):
            continue
        row_dict = {}
        for idx, val in enumerate(r):
            if idx < len(headers):
                row_dict[headers[idx]] = str(val).strip()
        parsed_rows.append(row_dict)
    return parsed_rows


def parse_excel_content(file_bytes):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [normalize_column_name(h) for h in rows[0] if h is not None]
        parsed_rows = []
        for r in rows[1:]:
            if not any(c is not None and str(c).strip() for c in r):
                continue
            row_dict = {}
            for idx, val in enumerate(r):
                if idx < len(headers):
                    if isinstance(val, datetime):
                        row_dict[headers[idx]] = val.strftime("%Y-%m-%d")
                    else:
                        row_dict[headers[idx]] = str(val).strip() if val is not None else ""
            parsed_rows.append(row_dict)
        return parsed_rows
    except ImportError:
        return parse_csv_content(file_bytes)


def parse_offer_file(file_bytes, filename):
    ext = Path(filename).suffix.lower()
    if ext not in (".csv", ".xlsx", ".xls"):
        return {
            "success": False,
            "error": "Unsupported file format. Please upload .csv, .xlsx, or .xls file."
        }

    try:
        if ext in (".xlsx", ".xls"):
            raw_rows = parse_excel_content(file_bytes)
        else:
            raw_rows = parse_csv_content(file_bytes)
    except Exception as exc:
        return {
            "success": False,
            "error": f"Failed to read file: {str(exc)}"
        }

    if not raw_rows:
        return {
            "success": False,
            "error": "The uploaded file contains no data rows."
        }

    candidates = []
    seen_emails_in_batch = set()

    for idx, r in enumerate(raw_rows, start=1):
        name = str(r.get("student_name") or "").strip()
        email_raw = str(r.get("student_email") or "").strip()
        if email_raw.endswith(".0") and "@" not in email_raw:
            email_raw = email_raw[:-2]
        email_clean = email_raw.lower()

        domain = str(r.get("domain") or "Full Stack Web Development").strip()
        raw_letter_type = str(r.get("letter_type") or "").strip()
        hours_per_week = str(r.get("hours_per_week") or "").strip()
        letter_type = normalize_letter_type(raw_letter_type, hours_per_week)

        raw_duration = r.get("duration")
        duration = normalize_duration(raw_duration)
        weeks = calculate_weeks_from_duration(duration)

        start_date_iso, start_date_display, start_err = parse_and_normalize_date(r.get("start_date"))
        end_date_iso, end_date_display, end_err = parse_and_normalize_date(r.get("end_date"))
        stipend = clean_stipend(r.get("stipend"))

        total_hours = calculate_total_hours(weeks, hours_per_week) if letter_type == "with_hours" else ""

        is_duplicate = False
        if email_clean:
            if email_clean in seen_emails_in_batch:
                is_duplicate = True
            else:
                seen_emails_in_batch.add(email_clean)

        # Validation checks
        errors = []
        if not name:
            errors.append("Name is required.")
        if not email_clean:
            errors.append("Email is required.")
        elif not is_valid_email(email_clean):
            errors.append("Invalid email format.")
        if not domain:
            errors.append("Domain is required.")
        if not start_date_iso:
            errors.append(start_err or "Start date is invalid.")
        if not end_date_iso:
            errors.append(end_err or "End date is invalid.")
        if start_date_iso and end_date_iso and end_date_iso < start_date_iso:
            errors.append("End date cannot be before start date.")
        if letter_type == "with_hours":
            if not hours_per_week:
                errors.append("Hours per week is required for With Hours.")
            else:
                try:
                    hpw = int(hours_per_week)
                    if hpw <= 0:
                        errors.append("Hours per week must be > 0.")
                except ValueError:
                    errors.append("Hours per week must be a number.")

        is_valid = len(errors) == 0

        candidates.append({
            "row_id": idx,
            "student_name": name,
            "student_email": email_clean,
            "domain": domain,
            "letter_type": letter_type,
            "duration": duration,
            "weeks": weeks,
            "hours_per_week": hours_per_week,
            "total_hours": total_hours or "",
            "start_date": start_date_iso or (r.get("start_date") or ""),
            "end_date": end_date_iso or (r.get("end_date") or ""),
            "stipend": stipend,
            "is_duplicate": is_duplicate,
            "is_valid": is_valid,
            "errors": errors,
            "status": "ready" if is_valid else "invalid",
        })

    return {
        "success": True,
        "total_rows": len(candidates),
        "candidates": candidates
    }


# ============================================================
# SAMPLE TEMPLATES
# ============================================================

def get_sample_csv_template():
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Candidate Name",
        "Email ID",
        "Internship Domain",
        "Letter Type",
        "Duration",
        "Start Date",
        "End Date",
        "Hours Per Week",
        "Stipend",
    ])
    writer.writerow([
        "Rahul Sharma",
        "rahul.sharma@example.com",
        "Full Stack Web Development",
        "With Hours",
        "2 months",
        "2026-06-01",
        "2026-07-31",
        "20",
        "5000",
    ])
    writer.writerow([
        "Ananya Patel",
        "ananya.patel@example.com",
        "Data Analytics",
        "Without Hours",
        "3 months",
        "2026-06-15",
        "2026-09-15",
        "",
        "",
    ])
    writer.writerow([
        "Vikram Verma",
        "vikram.verma@example.com",
        "Artificial Intelligence",
        "With Hours",
        "4 months",
        "2026-07-01",
        "2026-10-31",
        "25",
        "8000",
    ])
    return output.getvalue().encode("utf-8")


def get_sample_excel_template():
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Offer Letters"
        headers = [
            "Candidate Name", "Email ID", "Internship Domain", "Letter Type",
            "Duration", "Start Date", "End Date", "Hours Per Week", "Stipend"
        ]
        ws.append(headers)
        ws.append(["Rahul Sharma", "rahul.sharma@example.com", "Full Stack Web Development", "With Hours", "2 months", "2026-06-01", "2026-07-31", 20, 5000])
        ws.append(["Ananya Patel", "ananya.patel@example.com", "Data Analytics", "Without Hours", "3 months", "2026-06-15", "2026-09-15", "", ""])
        ws.append(["Vikram Verma", "vikram.verma@example.com", "Artificial Intelligence", "With Hours", "4 months", "2026-07-01", "2026-10-31", 25, 8000])
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    except ImportError:
        return get_sample_csv_template()
