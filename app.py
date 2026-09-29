import tempfile
import os
import io
import zipfile
import csv
import hmac
from pathlib import Path
import re
import smtplib
from datetime import datetime, timezone
from email.message import EmailMessage
import imaplib
import urllib.parse
import sqlite3
import uuid
from datetime import timedelta

from environment_config import load_application_environment, require_environment_variable, get_environment_info
from supabase import create_client, Client

APP_ENV = load_application_environment(Path(__file__).resolve().parent)
APP_LOCAL_TIMEZONE = timezone(timedelta(hours=5, minutes=30))

import sys
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from flask import (
    Flask,
    render_template,
    request,
    jsonify,
    send_from_directory,
    redirect,
    url_for,
    session,
    Response,
    send_file,
)

import pymupdf
import certificate_service
import database.repository as repository
import services.bulk_certificate_service as bulk_certificate_service
from services.ca_bulk_service import parse_ca_file
import services.bulk_offer_letter_service as bulk_offer_letter_service
from services import ca_certificate_service


# ============================================================
# APPLICATION
# ============================================================

app = Flask(__name__)

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:
    raise ValueError(
        "SUPABASE_URL and SUPABASE_KEY must be configured."
    )

supabase: Client = create_client(SUPABASE_URL,SUPABASE_KEY)


# ============================================================
# LOGIN / SESSION CONFIGURATION
# ============================================================
LOGIN_USERNAME = os.getenv("PERSEVEX_LOGIN_USERNAME", "").strip()
LOGIN_PASSWORD = os.getenv("PERSEVEX_LOGIN_PASSWORD", "").strip()
SESSION_SECRET = require_environment_variable("PERSEVEX_SESSION_SECRET")

app.secret_key = SESSION_SECRET
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = (
    os.getenv("VERCEL") == "1"
)


# ============================================================
# DIRECTORIES
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

TEMPLATE_DIR = BASE_DIR / "pdf_templates"

# Vercel allows writing only inside the temporary directory.
GENERATED_DIR = Path(tempfile.gettempdir()) / "persevex_generated"
FONT_DIR = Path(tempfile.gettempdir()) / "persevex_fonts"

GENERATED_DIR.mkdir(
    parents=True,
    exist_ok=True
)

FONT_DIR.mkdir(
    parents=True,
    exist_ok=True
)

# ============================================================
# EMAIL CONFIGURATION
# ============================================================

SENDER_EMAIL = os.getenv("SENDER_EMAIL", "").strip()

SENDER_PASSWORD = os.getenv("SENDER_PASSWORD","").strip()

SMTP_HOST = "smtpout.secureserver.net"
SMTP_PORT = 465

IMAP_HOST = "imap.secureserver.net"
IMAP_PORT = 993
IMAP_SENT_FOLDER = "Sent"


# ============================================================
# DURATION
#
# 1 MONTH IS REMOVED
# ============================================================

DURATION_WEEKS = {
    "2 months": 8,
    "3 months": 12,
    "4 months": 16,
    "5 months": 20,
    "6 months": 24,
}


# ============================================================
# BASIC HELPERS
# ============================================================

def safe_filename(name):

    name = re.sub(
        r"[^\w\s-]",
        "",
        str(name),
        flags=re.UNICODE,
    )

    name = name.strip()

    name = re.sub(
        r"\s+",
        "_",
        name,
    )

    return name or "Student"


def normalize_text(value):

    value = str(value)

    value = value.replace(
        "\u00a0",
        " ",
    )

    value = re.sub(
        r"\s+",
        " ",
        value,
    )

    return value.strip().lower()


def clean_stipend(value):

    if value is None:
        return ""

    value = str(value).strip()

    value = value.replace(
        "₹",
        "",
    )

    value = value.replace(
        ",",
        "",
    )

    s_clean = value.strip()

    return s_clean


def format_date(value):

    return datetime.strptime(
        value,
        "%Y-%m-%d",
    ).strftime(
        "%d/%m/%Y"
    )


def normalize_duration(value):

    value = str(value).strip().lower()

    if not value:
        return ""

    if re.fullmatch(
        r"\d+",
        value,
    ):

        return f"{value} months"

    return value


# ============================================================
# FONT EXTRACTION
# ============================================================

def extract_font(
    pdf_path,
    font_name,
):

    output = (
        FONT_DIR
        /
        f"{font_name}.ttf"
    )

    if output.exists():
        return str(output)

    doc = pymupdf.open(
        pdf_path
    )

    try:

        fonts = doc[0].get_fonts(
            full=True
        )

        for font in fonts:

            pdf_font_name = font[3]

            if (
                pdf_font_name == font_name
                or
                pdf_font_name.endswith(
                    "+" + font_name
                )
            ):

                extracted = doc.extract_font(
                    font[0]
                )

                if not extracted:
                    continue

                if (
                    len(extracted) >= 4
                    and
                    extracted[1].lower()
                    in (
                        "ttf",
                        "otf",
                    )
                ):

                    output.write_bytes(
                        extracted[3]
                    )

                    return str(output)

    finally:

        doc.close()

    raise RuntimeError(
        f"Could not extract "
        f"{font_name} from "
        f"{pdf_path.name}"
    )


# ============================================================
# FIND WINDOWS UNICODE FONT SUPPORTING ₹
# ============================================================

def find_rupee_font():

    windows_fonts = Path(
        r"C:\Windows\Fonts"
    )

    if not windows_fonts.exists():

        raise RuntimeError(
            "Windows Fonts directory was not found."
        )

    # --------------------------------------------------------
    # Prefer fonts that normally contain the Indian Rupee
    # symbol.
    # --------------------------------------------------------

    preferred = [
        "Nirmala.ttf",
        "NirmalaUI.ttf",
        "NirmalaUI-Regular.ttf",
        "segoeui.ttf",
        "arial.ttf",
        "calibri.ttf",
    ]

    for filename in preferred:

        font_path = (
            windows_fonts
            /
            filename
        )

        if not font_path.exists():
            continue

        try:

            font = pymupdf.Font(
                fontfile=str(font_path)
            )

            # Check actual glyph support.
            if font.has_glyph(
                ord("₹")
            ):

                return str(
                    font_path
                )

        except Exception:

            continue

    # --------------------------------------------------------
    # Search every installed font as fallback.
    # --------------------------------------------------------

    for font_path in windows_fonts.iterdir():

        if not font_path.is_file():
            continue

        if font_path.suffix.lower() not in (
            ".ttf",
            ".otf",
            ".ttc",
        ):
            continue

        try:

            font = pymupdf.Font(
                fontfile=str(font_path)
            )

            if font.has_glyph(
                ord("₹")
            ):

                return str(
                    font_path
                )

        except Exception:

            continue

    raise RuntimeError(
        "No installed Windows font supporting "
        "the ₹ symbol was found."
    )


# ============================================================
# GET TEXT LINES
# ============================================================

def get_text_lines(page):

    result = []

    blocks = page.get_text(
        "dict"
    ).get(
        "blocks",
        []
    )

    for block in blocks:

        if block.get(
            "type"
        ) != 0:

            continue

        for line in block.get(
            "lines",
            []
        ):

            spans = line.get(
                "spans",
                []
            )

            if not spans:
                continue

            text = "".join(
                span.get(
                    "text",
                    "",
                )
                for span in spans
            )

            if not text.strip():
                continue

            first = spans[0]

            result.append({

                "text": text,

                "bbox": pymupdf.Rect(
                    line["bbox"]
                ),

                "origin": first[
                    "origin"
                ],

                "font": first[
                    "font"
                ],

                "size": first[
                    "size"
                ],

                "color": first[
                    "color"
                ],

            })

    return result


# ============================================================
# FIND LINE
# ============================================================

def find_line(
    page,
    search_text,
):

    target = normalize_text(
        search_text
    )

    lines = get_text_lines(
        page
    )

    # Exact search.
    for line in lines:

        current = normalize_text(
            line["text"]
        )

        if target in current:

            return line

    # Word-based fallback.
    target_words = [
        word
        for word in re.findall(
            r"[a-zA-Z0-9]+",
            target,
        )
        if len(word) > 3
    ]

    for line in lines:

        current = normalize_text(
            line["text"]
        )

        matches = sum(
            1
            for word in target_words
            if word in current
        )

        if matches >= min(
            3,
            len(target_words),
        ):

            return line

    raise ValueError(
        f"Could not find PDF line: "
        f"{search_text}"
    )


def find_line_any(
    page,
    options,
):

    for option in options:

        try:

            return find_line(
                page,
                option,
            )

        except ValueError:

            pass

    raise ValueError(
        "Could not find PDF line: "
        +
        " / ".join(options)
    )


# ============================================================
# COLOR
# ============================================================

def rgb_from_pdf_color(
    color_value,
):

    r = (
        color_value >> 16
    ) & 255

    g = (
        color_value >> 8
    ) & 255

    b = color_value & 255

    return (
        r / 255,
        g / 255,
        b / 255,
    )


# ============================================================
# FIT FONT SIZE
# ============================================================

def fit_font_size(
    font_file,
    text,
    original_size,
    max_width,
):

    font = pymupdf.Font(
        fontfile=font_file
    )

    size = float(
        original_size
    )

    while size > 7:

        width = font.text_length(
            text,
            fontsize=size,
        )

        if width <= max_width:
            break

        size -= 0.1

    return size


# ============================================================
# REPLACE COMPLETE LINE
# ============================================================

def replace_entire_line(
    page,
    line,
    new_text,
    font_file,
):

    rect = pymupdf.Rect(
        line["bbox"]
    )

    original_size = float(
        line["size"]
    )

    color = rgb_from_pdf_color(
        line["color"]
    )

    fontsize = fit_font_size(
        font_file,
        new_text,
        original_size,
        rect.width,
    )

    page.add_redact_annot(
        rect,
        fill=(1, 1, 1),
    )

    page.apply_redactions()

    page.insert_text(
        (
            line["origin"][0],
            line["origin"][1],
        ),
        new_text,
        fontfile=font_file,
        fontsize=fontsize,
        color=color,
        overlay=True,
    )


# ============================================================
# REPLACE STUDENT NAME - KEEP ORIGINAL FONT SIZE
# ============================================================

def replace_student_name_line(
    page,
    line,
    new_text,
    font_file,
):
    """
    Replace the student-name line without reducing the template font size.

    Long names are wrapped onto the available space at the SAME font size.
    This helper is used only for the student-name line, so no other PDF
    formatting or replacement logic is changed.
    """

    rect = pymupdf.Rect(line["bbox"])
    original_size = float(line["size"])
    color = rgb_from_pdf_color(line["color"])

    # Remove only the original template text line.
    page.add_redact_annot(
        rect,
        fill=(1, 1, 1),
    )
    page.apply_redactions()

    # Keep the original left position and font size. Use the page width so
    # long names wrap instead of triggering the font-size reduction used by
    # replace_entire_line().
    right_margin = 30
    available = pymupdf.Rect(
        line["origin"][0],
        rect.y0 - 1,
        page.rect.width - right_margin,
        rect.y0 + max(90, original_size * 4.5),
    )

    result = page.insert_textbox(
        available,
        new_text,
        fontfile=font_file,
        fontsize=original_size,
        color=color,
        align=pymupdf.TEXT_ALIGN_LEFT,
        lineheight=1.0,
        overlay=True,
    )

    if result < 0:
        # Never shrink the name font. If an unusually long name needs more
        # vertical room, continue with the same font size in a larger box.
        available.y1 = page.rect.height - 30
        page.insert_textbox(
            available,
            new_text,
            fontfile=font_file,
            fontsize=original_size,
            color=color,
            align=pymupdf.TEXT_ALIGN_LEFT,
            lineheight=1.0,
            overlay=True,
        )


# ============================================================
# REPLACE WITH-HOURS OFFER LINE - KEEP ORIGINAL FONT SIZE
# ============================================================

def replace_hours_offer_line(
    page,
    line,
    new_text,
    font_file,
):
    """
    Replace ONLY the WITH HOURS offer line without reducing its font size.

    This function is intentionally limited to the line beginning
    "Has been offered an internship...". All other replacement logic remains
    unchanged.
    """

    rect = pymupdf.Rect(line["bbox"])
    original_size = float(line["size"])
    color = rgb_from_pdf_color(line["color"])

    # Clear only the original offer line.
    page.add_redact_annot(
        rect,
        fill=(1, 1, 1),
    )
    page.apply_redactions()

    # Use the available page width and keep the exact original font size.
    # The font size is never reduced.
    text_rect = pymupdf.Rect(
        line["origin"][0],
        rect.y0 - 1,
        page.rect.width - 24,
        rect.y1 + max(8, original_size * 1.5),
    )

    result = page.insert_textbox(
        text_rect,
        new_text,
        fontfile=font_file,
        fontsize=original_size,
        color=color,
        align=pymupdf.TEXT_ALIGN_LEFT,
        lineheight=1.0,
        overlay=True,
    )

    # If the text needs slightly more vertical room, retry without changing
    # the font size.
    if result < 0:
        text_rect.y1 = rect.y0 + max(
            original_size * 3.0,
            rect.height * 2.5,
        )
        page.insert_textbox(
            text_rect,
            new_text,
            fontfile=font_file,
            fontsize=original_size,
            color=color,
            align=pymupdf.TEXT_ALIGN_LEFT,
            lineheight=1.0,
            overlay=True,
        )


# ============================================================
# GET PDF WORDS
# ============================================================

def get_pdf_words(page):

    words = page.get_text(
        "words"
    )

    words.sort(
        key=lambda item: (
            item[5],
            item[6],
            item[7],
        )
    )

    return words


# ============================================================
# FIND STIPEND SENTENCE
# ============================================================

def find_stipend_sentence_words(
    page,
):

    words = get_pdf_words(
        page
    )

    if not words:
        return None

    cleaned = [
        normalize_text(
            word[4]
        )
        for word in words
    ]

    # --------------------------------------------------------
    # Find "The internship"
    # --------------------------------------------------------

    start_index = None

    for i in range(
        len(cleaned) - 1
    ):

        if (
            cleaned[i] == "the"
            and
            cleaned[i + 1]
            == "internship"
        ):

            nearby = " ".join(
                cleaned[
                    i:
                    min(
                        i + 25,
                        len(cleaned),
                    )
                ]
            )

            if (
                "stipend" in nearby
                or
                "performance-based"
                in nearby
                or
                "performance"
                in nearby
            ):

                start_index = i
                break

    # --------------------------------------------------------
    # Fallback
    # --------------------------------------------------------

    if start_index is None:

        for i in range(
            len(cleaned)
        ):

            nearby = " ".join(
                cleaned[
                    i:
                    min(
                        i + 20,
                        len(cleaned),
                    )
                ]
            )

            if (
                "performance-based"
                in nearby
                and
                "stipend"
                in nearby
            ):

                start_index = i

                for j in range(
                    max(
                        0,
                        i - 8,
                    ),
                    i,
                ):

                    if (
                        cleaned[j]
                        == "the"
                        and
                        j + 1
                        <
                        len(cleaned)
                        and
                        cleaned[j + 1]
                        == "internship"
                    ):

                        start_index = j
                        break

                break

    if start_index is None:
        return None

    # --------------------------------------------------------
    # Find stipend
    # --------------------------------------------------------

    stipend_index = None

    for i in range(
        start_index,
        min(
            start_index + 25,
            len(cleaned),
        ),
    ):

        if "stipend" in cleaned[i]:

            stipend_index = i
            break

    if stipend_index is None:
        return None

    # --------------------------------------------------------
    # Find "to"
    # --------------------------------------------------------

    amount_start = None

    for i in range(
        stipend_index,
        min(
            stipend_index + 15,
            len(cleaned),
        ),
    ):

        if cleaned[i] == "to":

            amount_start = i + 1
            break

    # --------------------------------------------------------
    # Find amount
    # --------------------------------------------------------

    amount_end = None

    if amount_start is not None:

        for i in range(
            amount_start,
            min(
                amount_start + 8,
                len(cleaned),
            ),
        ):

            raw = str(
                words[i][4]
            )

            value = cleaned[i]

            if "₹" in raw:

                if re.search(
                    r"\d",
                    raw,
                ):

                    amount_end = i
                    break

            if re.fullmatch(
                r"[\d,]+\.?",
                value,
            ):

                amount_end = i
                break

            if re.fullmatch(
                r"\d+\.?",
                value,
            ):

                amount_end = i
                break

    # --------------------------------------------------------
    # Find period
    # --------------------------------------------------------

    if amount_end is None:

        for i in range(
            stipend_index,
            min(
                stipend_index + 25,
                len(words),
            ),
        ):

            raw = str(
                words[i][4]
            )

            if (
                raw.endswith(".")
                or
                cleaned[i] == "."
            ):

                amount_end = i
                break

    # --------------------------------------------------------
    # Fallback
    # --------------------------------------------------------

    if amount_end is None:

        start_block = words[
            start_index
        ][5]

        amount_end = stipend_index

        for i in range(
            stipend_index + 1,
            len(words),
        ):

            current = words[i]

            if current[5] != start_block:
                break

            amount_end = i

            raw = str(
                current[4]
            )

            if raw.endswith("."):
                break

    if amount_end < start_index:
        amount_end = start_index

    return (
        start_index,
        amount_end,
        words,
    )


# ============================================================
# REMOVE COMPLETE STIPEND SENTENCE
# ============================================================

def remove_stipend_sentence(
    page,
):

    result = find_stipend_sentence_words(
        page
    )

    if not result:

        print(
            "STIPEND REMOVAL: "
            "Stipend sentence not found."
        )

        return

    start_index, end_index, words = result

    selected = words[
        start_index:
        end_index + 1
    ]

    if not selected:
        return

    line_groups = {}

    for word in selected:

        key = (
            word[5],
            word[6],
        )

        if key not in line_groups:

            line_groups[key] = []

        line_groups[key].append(
            word
        )

    rectangles = []

    for key, line_words in line_groups.items():

        if not line_words:
            continue

        x0 = min(
            word[0]
            for word in line_words
        )

        y0 = min(
            word[1]
            for word in line_words
        )

        x1 = max(
            word[2]
            for word in line_words
        )

        y1 = max(
            word[3]
            for word in line_words
        )

        rect = pymupdf.Rect(
            x0 - 2,
            y0 - 1,
            x1 + 2,
            y1 + 1,
        )

        rectangles.append(
            rect
        )

    for rect in rectangles:

        page.add_redact_annot(
            rect,
            fill=(1, 1, 1),
        )

    page.apply_redactions()

    print(
        "STIPEND REMOVAL: "
        "Complete stipend sentence removed."
    )


# ============================================================
# FIND STIPEND LINE
# ============================================================

def find_stipend_line(
    page,
):

    options = [
        "The internship offers a performance-based stipend",
        "internship offers a performance-based stipend",
        "offers a performance-based stipend",
        "performance-based stipend",
        "stipend of up to",
    ]

    return find_line_any(
        page,
        options,
    )


# ============================================================
# REPLACE STIPEND
#
# THIS VERSION FIXES ₹.
# ============================================================

def replace_stipend(
    page,
    template,
    stipend,
):
    """
    WITH HOURS ONLY: change ONLY the stipend amount.

    The existing sentence, paragraph layout, font, font size, spacing,
    and original ₹ symbol remain untouched.
    """

    stipend = clean_stipend(stipend)

    if not stipend:
        return

    line = find_stipend_line(page)

    # Keep the amount formatting clean while preserving the exact template text.
    try:
        formatted_amount = f"{int(float(stipend)):,}"
    except (TypeError, ValueError):
        formatted_amount = str(stipend)

    # --------------------------------------------------------
    # Find the amount span on the ORIGINAL stipend line.
    #
    # The Persevex template may store ₹ in a different font span
    # from the numeric amount. Therefore, replace only the number
    # whenever possible and leave the original ₹ glyph untouched.
    # --------------------------------------------------------

    amount_span = None
    line_y = line["bbox"].y0

    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue

        for pdf_line in block.get("lines", []):
            if abs(pdf_line["bbox"][1] - line_y) > 2.0:
                continue

            for span in pdf_line.get("spans", []):
                span_text = str(span.get("text", "")).strip()

                # Preferred case: amount is its own span, for example
                # "15,000." or "15000.".
                if re.fullmatch(r"[\d,]+\.?", span_text):
                    amount_span = span
                    break

                # Fallback: ₹ and amount are stored together.
                if "₹" in span_text and re.search(r"\d", span_text):
                    amount_span = span
                    break

            if amount_span is not None:
                break

        if amount_span is not None:
            break

    if amount_span is None:
        raise ValueError(
            "Could not locate the original stipend amount in the PDF template."
        )

    amount_rect = pymupdf.Rect(amount_span["bbox"])
    amount_text = str(amount_span.get("text", ""))

    # Remove ONLY the original amount. Do not remove or redraw the
    # surrounding stipend sentence.
    page.add_redact_annot(
        amount_rect,
        fill=(1, 1, 1),
    )
    page.apply_redactions()

    # --------------------------------------------------------
    # Normal template case:
    # ₹ is already present in its original font, so redraw ONLY
    # the number with the same typography as the original amount.
    # --------------------------------------------------------

    if "₹" not in amount_text:

        span_font_name = str(
            amount_span.get("font", "")
        )

        try:
            amount_font = extract_font(
                template,
                span_font_name,
            )
        except Exception:
            amount_font = extract_font(
                template,
                "OpenSans-Regular",
            )

        page.insert_text(
            amount_span["origin"],
            formatted_amount + ".",
            fontfile=amount_font,
            fontsize=float(amount_span["size"]),
            color=rgb_from_pdf_color(
                amount_span["color"]
            ),
            overlay=True,
        )

        print(
            "STIPEND: Amount updated. Original ₹ symbol and formatting preserved."
        )
        return

    # --------------------------------------------------------
    # Rare fallback:
    # ₹ and the amount are stored in one span. Preserve the same
    # font whenever it supports ₹; otherwise use the template's
    # embedded NotoSans font for the combined replacement.
    # --------------------------------------------------------

    span_font_name = str(
        amount_span.get("font", "")
    )

    try:
        combined_font = extract_font(
            template,
            span_font_name,
        )
        test_font = pymupdf.Font(
            fontfile=combined_font
        )

        if not test_font.has_glyph(
            ord("₹")
        ):
            raise RuntimeError(
                "Original amount font does not support ₹."
            )

    except Exception:

        try:
            combined_font = extract_font(
                template,
                "NotoSans-Regular",
            )
        except Exception:
            combined_font = find_rupee_font()

    page.insert_text(
        amount_span["origin"],
        f"₹{formatted_amount}.",
        fontfile=combined_font,
        fontsize=float(amount_span["size"]),
        color=rgb_from_pdf_color(
            amount_span["color"]
        ),
        overlay=True,
    )

    print(
        "STIPEND: Amount updated with ₹ while preserving the original layout."
    )


# ============================================================
# REPLACE DATE LINE
# ============================================================

def replace_date_line(
    page,
    template,
    start_date,
    end_date,
):

    line = find_line_any(
        page,
        [
            "The duration of the internship program ranges from",
            "duration of the internship program ranges from",
        ],
    )

    regular_font = extract_font(
        template,
        "OpenSans-Regular",
    )

    new_text = (
        "The duration of the internship program ranges from "
        f"{start_date} to {end_date} or"
    )

    replace_entire_line(
        page,
        line,
        new_text,
        regular_font,
    )


# ============================================================
# WITH HOURS
# ============================================================

def edit_with_hours(
    data,
):

    template = (
        TEMPLATE_DIR
        /
        "with_hours.pdf"
    )

    if not template.exists():

        raise FileNotFoundError(
            "WITH HOURS PDF template not found:\n"
            f"{template}"
        )

    doc = pymupdf.open(
        template
    )

    try:

        page = doc[0]

        regular_font = extract_font(
            template,
            "OpenSans-Regular",
        )

        # ----------------------------------------------------
        # STUDENT NAME
        # ----------------------------------------------------

        line = find_line_any(
            page,
            [
                "This letter is to confirm that Mr/Ms.",
                "This letter is to confirm",
            ],
        )

        new_text = (
            "This letter is to confirm that Mr/Ms. "
            f"{data['student_name']}"
        )

        replace_student_name_line(
            page,
            line,
            new_text,
            regular_font,
        )

        # ----------------------------------------------------
        # DOMAIN
        # ----------------------------------------------------

        line = find_line_any(
            page,
            [
                "Has been offered an internship in the field of",
                "Has been offered an internship",
            ],
        )

        new_text = (
            "Has been offered an internship in the field of "
            f"{data['domain']} with Persevex LLP,"
        )

        replace_hours_offer_line(
            page,
            line,
            new_text,
            regular_font,
        )

        # ----------------------------------------------------
        # DURATION
        # ----------------------------------------------------

        line = find_line_any(
            page,
            [
                "Which is of 2 months duration",
                "Which is of",
            ],
        )

        new_text = (
            "Which is of "
            f"{data['duration']} duration, "
            "under the supervision of "
            "Mr. Shanmukh Shekar K C."
        )

        replace_entire_line(
            page,
            line,
            new_text,
            regular_font,
        )

        # ----------------------------------------------------
        # HOURS FIRST LINE
        # ----------------------------------------------------

        line = find_line_any(
            page,
            [
                "standard requirements. I affirm that this internship",
                "standard requirements",
            ],
        )

        new_text = (
            "standard requirements. I affirm that this internship "
            f"is a total of {data['weeks']} weeks and "
            f"{data['hours_per_week']}"
        )

        replace_entire_line(
            page,
            line,
            new_text,
            regular_font,
        )

        # ----------------------------------------------------
        # HOURS SECOND LINE
        # ----------------------------------------------------

        line = find_line_any(
            page,
            [
                "hours per week, which sums to",
                "hours per week",
            ],
        )

        new_text = (
            "hours per week, which sums to "
            f"{data['total_hours']} hours for the entire "
            "duration of each domain."
        )

        replace_entire_line(
            page,
            line,
            new_text,
            regular_font,
        )

        # ----------------------------------------------------
        # STIPEND
        # ----------------------------------------------------

        stipend = clean_stipend(
            data.get(
                "stipend",
                "",
            )
        )

        if stipend:

            replace_stipend(
                page,
                template,
                stipend,
            )

        else:

            remove_stipend_sentence(
                page
            )

        # ----------------------------------------------------
        # DATES
        # ----------------------------------------------------

        replace_date_line(
            page,
            template,
            data["start_date"],
            data["end_date"],
        )

        return doc

    except Exception:

        doc.close()

        raise


# ============================================================
# WITHOUT-HOURS ONLY HELPERS
#
# IMPORTANT:
# These functions are intentionally separate from the WITH HOURS
# logic. The WITH HOURS section above is not changed.
# ============================================================

def find_without_hours_duration_block(page):
    """Find the original four-line training/internship paragraph."""

    lines = get_text_lines(page)
    start = None

    for line in lines:
        if "the total duration of the program is" in normalize_text(line["text"]):
            start = line
            break

    if start is None:
        raise ValueError(
            "Could not locate the without-hours duration paragraph."
        )

    # Collect lines immediately below the start line with the same left edge.
    candidates = [
        line for line in lines
        if line["bbox"].y0 >= start["bbox"].y0 - 0.5
        and line["bbox"].y0 <= start["bbox"].y0 + 55
        and abs(line["bbox"].x0 - start["bbox"].x0) < 2.0
    ]
    candidates.sort(key=lambda item: item["bbox"].y0)

    # The original template paragraph is exactly four visual lines.
    if len(candidates) < 4:
        raise ValueError(
            "Could not locate all four lines of the without-hours duration paragraph."
        )

    return candidates[:4]


def _replace_original_lines(page, lines, new_texts, font_file):
    """Redraw existing lines at their original positions and typography."""

    for line in lines:
        page.add_redact_annot(
            pymupdf.Rect(line["bbox"]),
            fill=(1, 1, 1),
        )

    page.apply_redactions()

    for line, text in zip(lines, new_texts):
        page.insert_text(
            line["origin"],
            text,
            fontfile=font_file,
            fontsize=line["size"],
            color=rgb_from_pdf_color(line["color"]),
            overlay=True,
        )


def replace_without_hours_duration_values(page, template, total_months):
    """
    WITHOUT HOURS ONLY.

    Preserve the original four-line paragraph exactly. Only the total-duration
    and internship-duration values are changed. This avoids the extra spaces
    caused by replacing individual words at different widths.
    """

    total_months = int(total_months)
    internship_months = total_months - 1

    total_unit = "month" if total_months == 1 else "months"
    internship_unit = "month" if internship_months == 1 else "months"

    lines = find_without_hours_duration_block(page)
    regular_font = extract_font(template, "OpenSans-Regular")

    original = [str(line["text"]) for line in lines]

    # First line: replace only "3 months" while preserving every other word.
    first_text = re.sub(
        r"\b3\s+months\b",
        f"{total_months} {total_unit}",
        original[0],
        count=1,
    )

    # Second line: replace only "2 months" while preserving every other word.
    second_text = re.sub(
        r"\b2\s+months\b",
        f"{internship_months} {internship_unit}",
        original[1],
        count=1,
    )

    new_texts = [
        first_text,
        second_text,
        original[2],
        original[3],
    ]

    _replace_original_lines(
        page,
        lines,
        new_texts,
        regular_font,
    )


def find_without_hours_stipend_parts(page):
    """Find the two visual parts of the wrapped stipend sentence."""

    lines = get_text_lines(page)
    previous = None
    stipend = None

    for line in lines:
        text = normalize_text(line["text"])

        if (
            "conduct weekly check-ins" in text
            and text.endswith("the")
        ):
            previous = line

        if (
            "internship offers a performance-based stipend" in text
            or "internship offers a performance-based stipend up to" in text
        ):
            stipend = line

    if previous is None or stipend is None:
        raise ValueError(
            "Could not locate the complete wrapped stipend sentence."
        )

    return previous, stipend


def replace_without_hours_stipend_amount(page, template, stipend):
    """WITHOUT HOURS ONLY: replace ONLY the original numeric stipend amount."""

    stipend = clean_stipend(stipend)
    _, stipend_line = find_without_hours_stipend_parts(page)

    regular_font = extract_font(
        template,
        "OpenSans-Regular",
    )

    fontsize = stipend_line["size"]
    color = rgb_from_pdf_color(stipend_line["color"])
    try:
        formatted = f"{int(float(stipend)):,}."
    except (ValueError, TypeError):
        formatted = f"{stipend}."

    # --------------------------------------------------------
    # IMPORTANT:
    # The original Persevex template stores ₹ as NotoSans-Regular
    # and the numeric amount as OpenSans-Regular. Therefore we find
    # the NUMERIC SPAN, not the PDF word. This leaves the original
    # ₹ glyph completely untouched.
    # --------------------------------------------------------

    amount_span = None

    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue

        for pdf_line in block.get("lines", []):
            if abs(pdf_line["bbox"][1] - stipend_line["bbox"].y0) > 1.0:
                continue

            for span in pdf_line.get("spans", []):
                text = str(span.get("text", ""))

                # Normal template: span is exactly "25,000.".
                if re.fullmatch(r"[\d,]+\.?", text.strip()):
                    amount_span = span
                    break

                # Fallback only if ₹ and number happen to be in one span.
                if "₹" in text and re.search(r"\d", text):
                    amount_span = span
                    break

            if amount_span is not None:
                break

        if amount_span is not None:
            break

    if amount_span is None:
        raise ValueError(
            "Could not locate the original stipend amount span."
        )

    amount_rect = pymupdf.Rect(
        amount_span["bbox"]
    )

    page.add_redact_annot(
        amount_rect,
        fill=(1, 1, 1),
    )
    page.apply_redactions()

    # Normal template: ₹ is already preserved as the original NotoSans
    # glyph, so insert only the numeric amount in OpenSans.
    if "₹" not in str(amount_span.get("text", "")):

        page.insert_text(
            (
                amount_span["origin"][0],
                stipend_line["origin"][1],
            ),
            formatted,
            fontfile=regular_font,
            fontsize=fontsize,
            color=color,
            overlay=True,
        )

        return

    # Rare fallback: the original PDF combined ₹ and the amount into one
    # span. Use the embedded NotoSans-Regular from the same template rather
    # than a system font, so the PDF remains self-contained.
    try:
        rupee_font = extract_font(
            template,
            "NotoSans-Regular",
        )
    except Exception:
        rupee_font = find_rupee_font()

    page.insert_text(
        (
            amount_span["origin"][0],
            stipend_line["origin"][1],
        ),
        f"₹{formatted}",
        fontfile=rupee_font,
        fontsize=fontsize,
        color=color,
        overlay=True,
    )

def remove_without_hours_stipend_sentence(page):
    """WITHOUT HOURS ONLY: remove the complete wrapped stipend sentence."""

    previous_line, stipend_line = find_without_hours_stipend_parts(page)
    words = page.get_text("words")

    trailing_the = None
    for word in words:
        if abs(word[1] - previous_line["bbox"].y0) <= 1.0:
            if normalize_text(word[4]) == "the":
                trailing_the = word

    if trailing_the is None:
        raise ValueError(
            "Could not locate the trailing 'The' of the stipend sentence."
        )

    page.add_redact_annot(
        pymupdf.Rect(
            trailing_the[0],
            trailing_the[1],
            trailing_the[2],
            trailing_the[3],
        ),
        fill=(1, 1, 1),
    )

    page.add_redact_annot(
        pymupdf.Rect(stipend_line["bbox"]),
        fill=(1, 1, 1),
    )

    page.apply_redactions()


def find_without_hours_date_lines(page):
    """Find the original two-line date paragraph."""

    lines = get_text_lines(page)
    first = None
    second = None

    for line in lines:
        text = normalize_text(line["text"])

        if "the duration of the internship program ranges from" in text:
            first = line
        elif "until the student completes the required number of internship hours" in text:
            second = line

    if first is None or second is None:
        raise ValueError(
            "Could not find the complete date paragraph."
        )

    return first, second


def replace_without_hours_date_paragraph(
    page,
    template,
    start_date,
    end_date,
    shift_up=0.0,
):
    """Replace only the dates and optionally move the original date paragraph."""

    first, second = find_without_hours_date_lines(page)
    regular_font = extract_font(template, "OpenSans-Regular")

    original_first = str(first["text"])
    original_second = str(second["text"])

    # Replace the two original dates in ONE operation so the newly inserted
    # start date cannot accidentally be matched again as the second date.
    matches = list(
        re.finditer(
            r"\d{2}/\d{2}/\d{4}",
            original_first,
        )
    )

    if len(matches) < 2:
        raise ValueError(
            "Could not locate both dates in the original date line."
        )

    new_first = (
        original_first[:matches[0].start()]
        + start_date
        + original_first[matches[0].end():matches[1].start()]
        + end_date
        + original_first[matches[1].end():]
    )

    # Remove the original paragraph before drawing the replacement.
    page.add_redact_annot(
        pymupdf.Rect(first["bbox"]),
        fill=(1, 1, 1),
    )
    page.add_redact_annot(
        pymupdf.Rect(second["bbox"]),
        fill=(1, 1, 1),
    )
    page.apply_redactions()

    page.insert_text(
        (
            first["origin"][0],
            first["origin"][1] - shift_up,
        ),
        new_first,
        fontfile=regular_font,
        fontsize=first["size"],
        color=rgb_from_pdf_color(first["color"]),
        overlay=True,
    )

    page.insert_text(
        (
            second["origin"][0],
            second["origin"][1] - shift_up,
        ),
        original_second,
        fontfile=regular_font,
        fontsize=second["size"],
        color=rgb_from_pdf_color(second["color"]),
        overlay=True,
    )


def replace_without_hours_date_paragraph_shifted(
    page,
    template,
    start_date,
    end_date,
    shift_up,
):
    """Compatibility wrapper for the shifted date operation."""

    replace_without_hours_date_paragraph(
        page,
        template,
        start_date,
        end_date,
        shift_up=shift_up,
    )


def replace_without_hours_offer_paragraph(page, offer_line, new_text, font_file):
    """
    Replace ONLY the Dates-section offer paragraph.

    The font size is taken from the normal body paragraph below it, never
    reduced to fit the text. The text wraps naturally while keeping the same
    font family, font size, spacing and left alignment for short and long
    student/domain values.
    """
    lines = get_text_lines(page)

    # Use the existing normal body paragraph as the formatting reference.
    duration_line = None
    for candidate in lines:
        candidate_text = normalize_text(candidate["text"])
        if "the total duration of the program is" in candidate_text:
            duration_line = candidate
            break

    if duration_line is None:
        raise ValueError("Could not locate the Dates-section body formatting reference.")

    body_size = float(duration_line["size"])
    color = rgb_from_pdf_color(duration_line["color"])

    # Clear only the existing offer paragraph area: from the original
    # 'Program in the field...' line up to the next body paragraph.
    clear_top = pymupdf.Rect(offer_line["bbox"]).y0 - 2
    clear_bottom = pymupdf.Rect(duration_line["bbox"]).y0 - 8
    clear_rect = pymupdf.Rect(
        offer_line["origin"][0] - 2,
        clear_top,
        page.rect.width - 24,
        clear_bottom,
    )

    page.add_redact_annot(clear_rect, fill=(1, 1, 1))
    page.apply_redactions()

    text_box = pymupdf.Rect(
        offer_line["origin"][0],
        clear_top,
        page.rect.width - 24,
        clear_bottom,
    )

    result = page.insert_textbox(
        text_box,
        new_text,
        fontfile=font_file,
        fontsize=body_size,
        color=color,
        align=pymupdf.TEXT_ALIGN_LEFT,
        lineheight=1.0,
        overlay=True,
    )

    if result < 0:
        raise ValueError(
            "The Dates-section offer paragraph does not fit at the original "
            "font size. The font was intentionally not reduced."
        )


# ============================================================
# WITHOUT HOURS
#
# IMPORTANT: WITH HOURS SECTION ABOVE IS LEFT UNCHANGED.
# ============================================================

def edit_without_hours(data):
    """Generate the WITHOUT HOURS letter only. WITH HOURS is untouched."""

    template = TEMPLATE_DIR / "without_hours.pdf"

    if not template.exists():
        raise FileNotFoundError(
            "WITHOUT HOURS PDF template not found:\n"
            f"{template}"
        )

    doc = pymupdf.open(template)

    try:
        page = doc[0]
        regular_font = extract_font(template, "OpenSans-Regular")

        # ----------------------------------------------------
        # STUDENT NAME - original behavior
        # ----------------------------------------------------
        line = find_line_any(
            page,
            [
                "This letter is to confirm that Mr/Ms.",
                "This letter is to confirm",
            ],
        )

        # Keep the student-name line independent from the next paragraph.
        # This prevents a long name from forcing the complete sentence into
        # the same text box and changing the visual layout.
        replace_student_name_line(
            page,
            line,
            "This letter is to confirm that Mr/Ms. "
            f"{data['student_name']}",
            regular_font,
        )

        # ----------------------------------------------------
        # DOMAIN - original behavior
        # ----------------------------------------------------
        line = find_line_any(
            page,
            [
                "Program in the field of Digital Marketing",
                "Program in the field of",
            ],
        )

        replace_without_hours_offer_paragraph(
            page,
            line,
            "Has been offered a Training and Internship Program in the field of "
            f"{data['domain']} with Persevex, under the supervision of "
            "Mr. Shanmukh Shekar K C.",
            regular_font,
        )

        # ----------------------------------------------------
        # DURATION - ONLY VALUES ARE CHANGED.
        # ----------------------------------------------------
        total_months = int(
            str(data["duration"]).split()[0]
        )

        replace_without_hours_duration_values(
            page,
            template,
            total_months,
        )

        # ----------------------------------------------------
        # STIPEND
        # ----------------------------------------------------
        stipend = clean_stipend(
            data.get("stipend", "")
        )

        if stipend:
            # Change ONLY the original amount. Keep the original ₹.
            replace_without_hours_stipend_amount(
                page,
                template,
                stipend,
            )

            # Dates stay exactly in their original position.
            replace_without_hours_date_paragraph(
                page,
                template,
                data["start_date"],
                data["end_date"],
                shift_up=0,
            )

        else:
            # Remove the entire wrapped stipend sentence, including the
            # trailing "The" on the preceding line.
            previous_line, stipend_line = find_without_hours_stipend_parts(page)

            # Date paragraph is currently below TWO removed visual lines:
            #   1) the trailing "The"
            #   2) the stipend sentence line
            # Move the date paragraph to the old stipend-line position.
            date_first, date_second = find_without_hours_date_lines(page)

            # Keep a small blank space between the paragraph above and
            # the date paragraph when no stipend is entered.
            # Only the vertical position is adjusted; all existing fonts,
            # sizes, formatting and other sections remain unchanged.
            desired_gap = 10
            shift_up = (
                date_first["origin"][1]
                - stipend_line["origin"][1]
                - desired_gap
            )

            remove_without_hours_stipend_sentence(page)

            replace_without_hours_date_paragraph(
                page,
                template,
                data["start_date"],
                data["end_date"],
                shift_up=shift_up,
            )

        return doc

    except Exception:
        doc.close()
        raise


# ============================================================
# GENERATE PDF
# ============================================================

def generate_pdf(
    letter_type,
    data,
):

    if letter_type == "with_hours":

        doc = edit_with_hours(
            data
        )

    elif letter_type == "without_hours":

        doc = edit_without_hours(
            data
        )

    else:

        raise ValueError(
            "Invalid letter type."
        )

    filename = (
        f"{safe_filename(data['student_name'])}"
        "_Internship_Acceptance_Letter.pdf"
    )

    output = (
        GENERATED_DIR
        /
        filename
    )

    if output.exists():

        output.unlink()

    try:

        doc.save(
            output,
            garbage=4,
            deflate=True,
        )

    finally:

        doc.close()

    return filename


# ============================================================
# CAMPUS AMBASSADOR OFFER LETTER GENERATOR
# ============================================================

def format_ca_date(date_val):
    """
    Format any date representation into 'Month DD, YYYY' matching the CA template.
    Defaults to current date if missing or invalid.
    """
    if not date_val:
        return datetime.now().strftime("%B %d, %Y")
    date_val = str(date_val).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%B %d, %Y", "%d %B %Y", "%B %d %Y", "%Y/%m/%d"):
        try:
            dt = datetime.strptime(date_val, fmt)
            return dt.strftime("%B %d, %Y")
        except ValueError:
            pass
    return date_val


_CA_TEMPLATE_BYTES = None

def get_ca_template_bytes():
    global _CA_TEMPLATE_BYTES
    if _CA_TEMPLATE_BYTES is None:
        template = TEMPLATE_DIR / "campus_ambassador_template.pdf"
        if not template.exists():
            raise FileNotFoundError(
                f"Campus Ambassador PDF template not found:\n{template}"
            )
        _CA_TEMPLATE_BYTES = template.read_bytes()
    return _CA_TEMPLATE_BYTES


def edit_campus_ambassador(data):
    """
    Generate the Campus Ambassador Offer Letter.
    Preserves 2-page master template exactly.
    Redacts and replaces ONLY:
      1. Candidate Name (after 'Dear ')
      2. Date (below 'Date :')
    Page 2 is 100% untouched.
    """
    t_bytes = get_ca_template_bytes()
    doc = pymupdf.open(stream=t_bytes, filetype="pdf")
    try:
        page = doc[0]

        # Redact candidate name text area (after 'Dear ')
        name_rect = pymupdf.Rect(105.0, 220.0, 380.0, 252.0)
        page.add_redact_annot(name_rect, fill=(1, 1, 1))

        # Redact date text area (strictly below y=233.5 to preserve 'Date :')
        date_rect = pymupdf.Rect(410.0, 234.0, 555.0, 252.0)
        page.add_redact_annot(date_rect, fill=(1, 1, 1))

        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE)

        # Insert Candidate Name
        student_name = (data.get("student_name") or "").strip()
        page.insert_text(
            pymupdf.Point(108.0, 244.2),
            student_name,
            fontsize=18,
            fontname="helv",
            color=(0, 0, 0),
        )

        # Insert Date formatted as 'Month DD, YYYY'
        raw_date = data.get("date") or data.get("issue_date") or data.get("start_date")
        date_str = format_ca_date(raw_date)
        font_size = 14
        text_w = pymupdf.get_text_length(date_str, fontname="helv", fontsize=font_size)
        date_x = 528.0 - text_w
        page.insert_text(
            pymupdf.Point(date_x, 246.0),
            date_str,
            fontsize=font_size,
            fontname="helv",
            color=(0, 0, 0),
        )

        return doc
    except Exception:
        doc.close()
        raise


def generate_ca_pdf(data):
    doc = edit_campus_ambassador(data)
    student_name = (data.get("student_name") or "Student").strip()
    filename = f"{safe_filename(student_name)}_Campus_Ambassador_Offer_Letter.pdf"
    output = GENERATED_DIR / filename
    if output.exists():
        try:
            output.unlink()
        except OSError:
            pass

    try:
        doc.save(
            output,
            garbage=4,
            deflate=True,
        )
    finally:
        doc.close()

    return filename


# ============================================================
# LOGIN PROTECTION
# ============================================================

@app.before_request
def require_login():
    allowed_endpoints = {
        "login",
        "static",
        "verify_certificate",
        "generated",
    }

    if request.endpoint in allowed_endpoints:
        return None

    if request.path.startswith("/verify") or request.path.startswith("/static"):
        return None

    if session.get("authenticated") is True:
        return None

    if request.path.startswith("/generate") or request.path.startswith("/send-email") or request.path.startswith("/api/"):
        return jsonify({
            "error": "Login required. Please log in first."
        }), 401

    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("authenticated") is True:
        return redirect(url_for("index"))

    error = None

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        if not LOGIN_USERNAME or not LOGIN_PASSWORD:
            error = "Login is not configured. Set the login environment variables."
        elif (
            hmac.compare_digest(username, LOGIN_USERNAME)
            and hmac.compare_digest(password, LOGIN_PASSWORD)
        ):
            session.clear()
            session["authenticated"] = True
            session["username"] = LOGIN_USERNAME
            return redirect(url_for("index"))
        else:
            error = "Invalid username or password."

    return render_template("login.html", error=error)


@app.get("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ============================================================
# HOME
# ============================================================

@app.get("/")
def index():
    try:
        history_data = repository.get_unified_history(record_type="offer_letter", page=1, per_page=1)
        success_count = history_data.get("total_sent", 0)
        failed_count = history_data.get("total_failed", 0)
    except Exception as exc:
        print("HEADER STATUS COUNT ERROR:", repr(exc))
        success_count = 0
        failed_count = 0

    return render_template(
        "index.html",
        success_count=success_count,
        failed_count=failed_count,
    )

# ============================================================
# LIVE EMAIL STATUS COUNTS
# ============================================================

@app.get("/api/email-status-counts")
def email_status_counts():
    try:
        history_data = repository.get_unified_history(record_type="offer_letter", page=1, per_page=1)
        return jsonify({
            "success": True,
            "total_count": history_data.get("total_offer_letters", 0),
            "success_count": history_data.get("total_sent", 0),
            "failed_count": history_data.get("total_failed", 0),
        })
    except Exception as exc:
        print("LIVE STATUS COUNT ERROR:", repr(exc))
        return jsonify({
            "success": False,
            "total_count": 0,
            "success_count": 0,
            "failed_count": 0
        }), 500

# ============================================================
# GENERATE API
# ============================================================

@app.post("/generate")
def generate():

    try:

        data = request.get_json(
            force=True
        )

        if not data:

            return jsonify({
                "error": "No data received."
            }), 400

        letter_type = data.get(
            "letter_type"
        )

        if letter_type not in (
            "with_hours",
            "without_hours",
        ):

            return jsonify({
                "error": "Invalid letter type."
            }), 400

        # ----------------------------------------------------
        # REQUIRED FIELDS
        #
        # STIPEND IS OPTIONAL.
        # ----------------------------------------------------

        required = [
            "student_name",
            "student_email",
            "domain",
            "duration",
            "start_date",
            "end_date",
        ]

        if letter_type == "with_hours":

            required.append(
                "hours_per_week"
            )

        missing = []

        for field in required:

            value = data.get(
                field,
                "",
            )

            if not str(
                value
            ).strip():

                missing.append(
                    field
                )

        if missing:

            return jsonify({
                "error": (
                    "Please fill: "
                    +
                    ", ".join(
                        missing
                    )
                )
            }), 400

        # ----------------------------------------------------
        # DURATION
        # ----------------------------------------------------

        duration = normalize_duration(
            data["duration"]
        )

        if duration not in DURATION_WEEKS:

            return jsonify({
                "error": (
                    "Please select a duration "
                    "between 2 and 6 months."
                )
            }), 400

        data["duration"] = duration

        # ----------------------------------------------------
        # DATE VALIDATION
        # ----------------------------------------------------

        try:

            start_date_obj = datetime.strptime(
                data["start_date"],
                "%Y-%m-%d",
            )

            end_date_obj = datetime.strptime(
                data["end_date"],
                "%Y-%m-%d",
            )

        except ValueError:

            return jsonify({
                "error": "Invalid date format."
            }), 400

        if end_date_obj < start_date_obj:

            return jsonify({
                "error": (
                    "End date cannot be before "
                    "start date."
                )
            }), 400

        # ----------------------------------------------------
        # CLEAN DATA
        # ----------------------------------------------------

        data["student_name"] = str(
            data["student_name"]
        ).strip()

        data["student_email"] = str(
            data["student_email"]
        ).strip()

        data["domain"] = str(
            data["domain"]
        ).strip()

        data["start_date"] = format_date(
            data["start_date"]
        )

        data["end_date"] = format_date(
            data["end_date"]
        )

        # ----------------------------------------------------
        # OPTIONAL STIPEND
        # ----------------------------------------------------

        data["stipend"] = clean_stipend(
            data.get(
                "stipend",
                "",
            )
        )

        # ----------------------------------------------------
        # WITH HOURS
        # ----------------------------------------------------

        if letter_type == "with_hours":

            weeks = DURATION_WEEKS[
                duration
            ]

            try:

                hours_per_week = int(
                    data["hours_per_week"]
                )

            except (
                ValueError,
                TypeError,
            ):

                return jsonify({
                    "error": (
                        "Hours per week must "
                        "be a number."
                    )
                }), 400

            if hours_per_week <= 0:

                return jsonify({
                    "error": (
                        "Hours per week must "
                        "be greater than zero."
                    )
                }), 400

            total_hours = (
                weeks
                *
                hours_per_week
            )

            data["weeks"] = str(
                weeks
            )

            data["hours_per_week"] = str(
                hours_per_week
            )

            data["total_hours"] = str(
                total_hours
            )

        # ----------------------------------------------------
        # GENERATE
        # ----------------------------------------------------

        filename = generate_pdf(
            letter_type,
            data,
        )

        return jsonify({

            "success": True,

            "filename": filename,

            "url": (
                f"/generated/{filename}"
            ),

        })

    except Exception as exc:

        print(
            "========================================"
        )

        print(
            "PDF GENERATION ERROR:"
        )

        print(
            repr(exc)
        )

        print(
            "========================================"
        )

        return jsonify({
            "error": str(exc)
        }), 500


# ============================================================
# SERVE GENERATED PDF
# ============================================================

@app.get(
    "/generated/<path:filename>"
)
def generated(
    filename,
):

    return send_from_directory(
        GENERATED_DIR,
        filename,
        as_attachment=False,
    )

# ============================================================
# EMAIL HISTORY HELPERS
# ============================================================

def get_previous_email_record(email):

    try:

        response = (
            supabase
            .table("email_history")
            .select("*")
            .eq(
                "student_email",
                email
            )
            .order(
                "created_at",
                desc=True
            )
            .limit(1)
            .execute()
        )

        if response.data:

            return response.data[0]

        return None

    except Exception as exc:

        print(
            "DUPLICATE EMAIL CHECK ERROR:",
            repr(exc)
        )

        return None
# ============================================================
# SEND EMAIL
# ============================================================

@app.post("/send-email")
def send_email():

    data = request.get_json(force=True) or {}

    # --------------------------------------------------------
    # GET DATA
    # --------------------------------------------------------

    filename = data.get("filename") or ""

    recipient = (
        data.get("student_email") or ""
    ).strip().lower()

    student_name = (
        data.get("student_name") or "Student"
    ).strip()

    send_again = (
        data.get("send_again") is True
    )

    # --------------------------------------------------------
    # HELPER FUNCTION
    # --------------------------------------------------------

    def build_history_data(
        status,
        sent_at=None,
        send_count=0,
        error_message=None
    ):

        safe_filename = (
            Path(filename).name
            if filename
            else None
        )

        offer_letter_id = (
            data.get("offer_letter_id")
            or (
                Path(safe_filename).stem
                if safe_filename
                else None
            )
        )

        return {
            "student_name": student_name,
            "student_email": recipient,
            "phone_number": data.get("phone_number"),
            "college_name": data.get("college_name"),

            "internship_domain": (
                data.get("internship_domain")
                or data.get("domain")
            ),

            "internship_duration": (
                data.get("internship_duration")
                or data.get("duration")
            ),

            "start_date": data.get("start_date"),
            "end_date": data.get("end_date"),

            "offer_letter_type": (
                data.get("offer_letter_type")
                or data.get("letter_type")
            ),

            "email_status": status,
            "sent_at": sent_at,
            "offer_letter_id": offer_letter_id,
            "pdf_filename": safe_filename,
            "send_count": send_count,
            "error_message": error_message
        }

    # --------------------------------------------------------
    # VALIDATION
    # --------------------------------------------------------

    if not filename:
        return jsonify({
            "success": False,
            "error": "Generated PDF is required."
        }), 400

    if not recipient:
        return jsonify({
            "success": False,
            "error": "Student email is required."
        }), 400

    # --------------------------------------------------------
    # SAFE FILENAME
    # --------------------------------------------------------

    filename = Path(filename).name

    pdf_path = GENERATED_DIR / filename

    if not pdf_path.exists():
        return jsonify({
            "success": False,
            "error": "Generated PDF not found. Please generate the letter again."
        }), 404

    # --------------------------------------------------------
    # CHECK PREVIOUS EMAIL
    # --------------------------------------------------------

    previous_record = get_previous_email_record(recipient)

    # --------------------------------------------------------
    # DUPLICATE CHECK
    # --------------------------------------------------------

    if previous_record and not send_again:
        return jsonify({
            "success": False,
            "duplicate": True,
            "message": "An offer letter record already exists for this email address.",
            "previous_record": previous_record
        }), 409

    # --------------------------------------------------------
    # EMAIL CONFIGURATION
    # --------------------------------------------------------

    if not SENDER_EMAIL or not SENDER_PASSWORD:
        return jsonify({
            "success": False,
            "error": "Email is not configured."
        }), 500

    # ========================================================
    # PREPARE EMAIL
    # ========================================================

    try:

        message = EmailMessage()

        message["From"] = f"Persevex LLP <{SENDER_EMAIL}>"

        message["To"] = recipient

        message["Subject"] = "Internship Acceptance Letter"

        # ----------------------------------------------------
        # PLAIN TEXT EMAIL
        # ----------------------------------------------------

        message.set_content(
f"""Dear {student_name},

Warm greetings from Persevex LLP!

We are delighted to inform you that your Internship Acceptance Letter has been attached with this email. Please review the document carefully and feel free to reach out if you need any clarification.

We are truly excited to welcome you onboard at Persevex and look forward to your active contribution and learning journey with us. Your enthusiasm and dedication will play a key role in shaping meaningful experiences throughout this internship.

Kindly acknowledge the receipt of this email and confirm your acceptance at your earliest convenience.

Wishing you a wonderful start with us!

Warm regards,
Team Persevex"""
        )

        # ----------------------------------------------------
        # HTML EMAIL
        # ----------------------------------------------------

        message.add_alternative(
f"""
<!DOCTYPE html>
<html>
<body style="font-family: Arial, Helvetica, sans-serif; font-size: 16px; line-height: 1.6; color: #333333;">

<p>Dear {student_name},</p>

<p>Warm greetings from Persevex LLP!</p>

<p>
We are delighted to inform you that your
<strong>Internship Acceptance Letter</strong>
has been attached with this email. Please review the document carefully
and feel free to reach out if you need any clarification.
</p>

<p>
We are truly excited to welcome you onboard at Persevex and look forward
to your active contribution and learning journey with us. Your enthusiasm
and dedication will play a key role in shaping meaningful experiences
throughout this internship.
</p>

<p>
Kindly acknowledge the receipt of this email and confirm your acceptance
at your earliest convenience.
</p>

<p>Wishing you a wonderful start with us!</p>

<p>
Warm regards,<br>
<strong>Team Persevex</strong>
</p>

</body>
</html>
""",
            subtype="html"
        )

        # ----------------------------------------------------
        # ATTACH PDF
        # ----------------------------------------------------

        message.add_attachment(
            pdf_path.read_bytes(),
            maintype="application",
            subtype="pdf",
            filename=filename
        )

    except Exception as exc:

        print("EMAIL PREPARATION ERROR:", repr(exc))

        return jsonify({
            "success": False,
            "error": str(exc),
            "email_status": "failed"
        }), 500

    # ========================================================
    # SEND EMAIL USING SMTP
    # ========================================================

    try:

        print("========================================")
        print("CONNECTING TO SMTP")
        print("FROM:", SENDER_EMAIL)
        print("TO:", recipient)
        print("========================================")

        with smtplib.SMTP_SSL(
            SMTP_HOST,
            SMTP_PORT,
            timeout=30
        ) as smtp:

            smtp.login(
                SENDER_EMAIL,
                SENDER_PASSWORD
            )

            smtp.send_message(message)

        print("SMTP EMAIL SENT SUCCESSFULLY")

    except Exception as exc:

        error_message = str(exc)

        print("========================================")
        print("SMTP EMAIL FAILED")
        print("Error:", repr(exc))
        print("========================================")

        # ----------------------------------------------------
        # SAVE FAILED HISTORY
        # ----------------------------------------------------

        try:

            if previous_record:

                previous_id = previous_record.get("id")

                previous_send_count = int(
                    previous_record.get("send_count") or 0
                )

                failed_history_data = build_history_data(
                    status="failed",
                    sent_at=None,
                    send_count=previous_send_count + 1,
                    error_message=error_message
                )

                supabase.table(
                    "email_history"
                ).update(
                    failed_history_data
                ).eq(
                    "id",
                    previous_id
                ).execute()

            else:

                failed_history_data = build_history_data(
                    status="failed",
                    sent_at=None,
                    send_count=1,
                    error_message=error_message
                )

                supabase.table(
                    "email_history"
                ).insert(
                    failed_history_data
                ).execute()

        except Exception as supabase_error:

            print(
                "FAILED HISTORY SAVE ERROR:",
                repr(supabase_error)
            )

        return jsonify({
            "success": False,
            "error": error_message,
            "email_status": "failed"
        }), 500

    # ========================================================
    # EMAIL SENT SUCCESSFULLY
    # ========================================================

    sent_time = datetime.now(
        timezone.utc
    ).isoformat()

    # --------------------------------------------------------
    # SAVE SENT EMAIL TO IMAP SENT FOLDER
    # --------------------------------------------------------

    try:

        with imaplib.IMAP4_SSL(
            IMAP_HOST,
            993
        ) as imap:

            imap.login(
                SENDER_EMAIL,
                SENDER_PASSWORD
            )

            result = imap.append(
                IMAP_SENT_FOLDER,
                "\\Seen",
                None,
                message.as_bytes()
            )

            print(
                "EMAIL SAVED TO SENT FOLDER:",
                result
            )

    except Exception as imap_error:

        # Email was already sent successfully.
        # Do not mark the entire request as failed.
        print(
            "WARNING: Could not save email to Sent folder:",
            repr(imap_error)
        )

    # --------------------------------------------------------
    # RESEND EMAIL
    # --------------------------------------------------------

    if previous_record and send_again:

        try:

            previous_id = previous_record.get("id")

            previous_send_count = int(
                previous_record.get("send_count") or 0
            )

            update_data = build_history_data(
                status="sent",
                sent_at=sent_time,
                send_count=previous_send_count + 1,
                error_message=None
            )

            supabase.table(
                "email_history"
            ).update(
                update_data
            ).eq(
                "id",
                previous_id
            ).execute()

            print(
                "EMAIL RESEND HISTORY UPDATED SUCCESSFULLY"
            )

        except Exception as supabase_error:

            print(
                "SUPABASE RESEND UPDATE ERROR:",
                repr(supabase_error)
            )

        return jsonify({
            "success": True,
            "resent": True,
            "message": "Email sent again successfully."
        })

    # --------------------------------------------------------
    # NEW EMAIL - SAVE HISTORY
    # --------------------------------------------------------

    try:

        sent_history_data = build_history_data(
            status="sent",
            sent_at=sent_time,
            send_count=1,
            error_message=None
        )

        supabase.table(
            "email_history"
        ).insert(
            sent_history_data
        ).execute()

        print(
            "EMAIL HISTORY SAVED SUCCESSFULLY"
        )

    except Exception as supabase_error:

        print(
            "SUPABASE HISTORY SAVE ERROR:",
            repr(supabase_error)
        )

    # ========================================================
    # SUCCESS RESPONSE
    # ========================================================

    return jsonify({
        "success": True,
        "message": "Email sent successfully."
    })


# ============================================================
# BULK OFFER LETTER ROUTES
# ============================================================

@app.post("/api/offer-letter/bulk/parse-upload")
def bulk_offer_letter_parse_upload():
    """
    Parse uploaded CSV/Excel file or folder of files for bulk offer letters.
    """
    try:
        uploaded_files = request.files.getlist("files") or request.files.getlist("file")
        if not uploaded_files:
            file = request.files.get("file")
            if file:
                uploaded_files = [file]

        if not uploaded_files:
            return jsonify({"success": False, "error": "No files selected."}), 400

        all_candidates = []
        files_scanned = 0
        files_supported = 0
        files_skipped = 0
        skipped_names = []
        row_counter = 1

        for f in uploaded_files:
            if not f or not f.filename:
                continue
            files_scanned += 1
            ext = Path(f.filename).suffix.lower()
            if ext not in (".csv", ".xlsx", ".xls"):
                files_skipped += 1
                skipped_names.append(Path(f.filename).name)
                continue

            files_supported += 1
            file_bytes = f.read()
            res = bulk_offer_letter_service.parse_offer_file(file_bytes, f.filename)
            if res.get("success"):
                for c in res.get("candidates", []):
                    c["row_id"] = row_counter
                    row_counter += 1
                    all_candidates.append(c)

        if files_supported == 0:
            return jsonify({
                "success": False,
                "error": "No supported CSV or Excel files found. Supported formats: .csv, .xlsx, .xls",
                "files_scanned": files_scanned,
                "files_skipped": files_skipped,
                "skipped_files": skipped_names
            }), 400

        # Recalculate in-batch duplicates across the whole aggregated list
        seen_emails = set()
        duplicate_count = 0
        valid_count = 0
        with_hours_count = 0

        for c in all_candidates:
            em = str(c.get("student_email") or "").strip().lower()
            if em:
                if em in seen_emails:
                    c["is_duplicate"] = True
                    duplicate_count += 1
                else:
                    seen_emails.add(em)
                    c["is_duplicate"] = False
            if c.get("is_valid"):
                valid_count += 1
            if c.get("letter_type") == "with_hours":
                with_hours_count += 1

        return jsonify({
            "success": True,
            "total_rows": len(all_candidates),
            "valid_rows": valid_count,
            "invalid_rows": len(all_candidates) - valid_count,
            "duplicate_rows": duplicate_count,
            "with_hours_rows": with_hours_count,
            "candidates": all_candidates,
            "files_scanned": files_scanned,
            "files_supported": files_supported,
            "files_skipped": files_skipped,
            "skipped_files": skipped_names
        })

    except Exception as exc:
        print("BULK OFFER LETTER PARSE UPLOAD ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.post("/api/offer-letter/bulk/check-duplicates")
def bulk_offer_letter_check_duplicates():
    """
    Check batch emails against email_history for previously sent recipients.
    """
    try:
        data = request.get_json(force=True) or {}
        records = data.get("records") or data.get("emails") or []
        if not records:
            return jsonify({
                "success": True,
                "total_checked": 0,
                "duplicate_count": 0,
                "duplicates_count": 0,
                "duplicates": [],
                "existing_emails": [],
                "duplicate_map": {}
            })

        duplicate_map = repository.find_existing_offer_letters_batch(records, supabase_client=supabase)
        duplicate_emails = list(duplicate_map.keys())

        return jsonify({
            "success": True,
            "total_checked": len(records),
            "duplicate_count": len(duplicate_emails),
            "duplicates_count": len(duplicate_emails),
            "duplicates": duplicate_emails,
            "existing_emails": duplicate_emails,
            "duplicate_map": duplicate_map
        })
    except Exception as exc:
        print("BULK OFFER LETTER CHECK DUPLICATES ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.post("/api/offer-letter/bulk/generate-item")
def bulk_offer_letter_generate_item():
    """
    Generate a single Offer Letter PDF for bulk flow without dispatching email or creating history records.
    """
    try:
        data = request.get_json(force=True) or {}
        letter_type = data.get("letter_type") or "with_hours"
        if letter_type not in ("with_hours", "without_hours"):
            return jsonify({"success": False, "error": "Invalid letter type."}), 400

        student_name = str(data.get("student_name") or data.get("candidate_name") or "").strip()
        student_email = str(data.get("student_email") or data.get("candidate_email") or "").strip().lower()
        domain = str(data.get("domain") or data.get("internship_domain") or "").strip()
        duration_raw = data.get("duration") or "2 months"
        duration = normalize_duration(duration_raw)
        start_date = str(data.get("start_date") or "").strip()
        end_date = str(data.get("end_date") or "").strip()
        stipend = clean_stipend(data.get("stipend"))

        if not student_name or not student_email or not domain or not start_date or not end_date:
            return jsonify({"success": False, "error": "Missing required fields."}), 400

        # Format dates to DD/MM/YYYY for template
        try:
            if "-" in start_date and len(start_date.split("-")[0]) == 4:
                s_obj = datetime.strptime(start_date, "%Y-%m-%d")
                formatted_start = format_date(start_date)
            else:
                s_obj = datetime.strptime(start_date, "%d/%m/%Y")
                formatted_start = start_date

            if "-" in end_date and len(end_date.split("-")[0]) == 4:
                e_obj = datetime.strptime(end_date, "%Y-%m-%d")
                formatted_end = format_date(end_date)
            else:
                e_obj = datetime.strptime(end_date, "%d/%m/%Y")
                formatted_end = end_date
        except ValueError:
            return jsonify({"success": False, "error": "Invalid date format."}), 400

        if e_obj < s_obj:
            return jsonify({"success": False, "error": "End date cannot be before start date."}), 400

        data_payload = {
            "student_name": student_name,
            "student_email": student_email,
            "domain": domain,
            "duration": duration,
            "start_date": formatted_start,
            "end_date": formatted_end,
            "stipend": stipend,
        }

        if letter_type == "with_hours":
            weeks = DURATION_WEEKS.get(duration, 8)
            try:
                hours_per_week = int(data.get("hours_per_week") or 20)
            except (ValueError, TypeError):
                return jsonify({"success": False, "error": "Hours per week must be a number."}), 400
            if hours_per_week <= 0:
                return jsonify({"success": False, "error": "Hours per week must be greater than zero."}), 400

            total_hours = weeks * hours_per_week
            data_payload["weeks"] = str(weeks)
            data_payload["hours_per_week"] = str(hours_per_week)
            data_payload["total_hours"] = str(total_hours)

        filename = generate_pdf(letter_type, data_payload)

        return jsonify({
            "success": True,
            "filename": filename,
            "url": f"/generated/{filename}",
            "student_name": student_name,
            "student_email": student_email
        })

    except Exception as exc:
        print("BULK OFFER LETTER GENERATE ITEM ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.post("/api/offer-letter/bulk/send-item")
def bulk_offer_letter_send_item():
    """
    Send a single offer letter email and update history using existing offer letter persistence rules.
    """
    try:
        data = request.get_json(force=True) or {}
        filename = (data.get("filename") or "").strip()
        recipient = (data.get("student_email") or "").strip().lower()
        student_name = (data.get("student_name") or "Student").strip()
        domain = (data.get("domain") or data.get("internship_domain") or "").strip()
        duration = normalize_duration(data.get("duration") or "2 months")
        start_date = str(data.get("start_date") or "").strip()
        end_date = str(data.get("end_date") or "").strip()
        letter_type = data.get("letter_type") or data.get("offer_letter_type") or "with_hours"
        stipend = clean_stipend(data.get("stipend"))
        send_again = bool(data.get("send_again", False))

        if not recipient:
            return jsonify({"success": False, "error": "Student email is required."}), 400

        # Generate PDF on demand if not present on disk
        if not filename or not (GENERATED_DIR / Path(filename).name).exists():
            if "-" in start_date and len(start_date.split("-")[0]) == 4:
                formatted_start = format_date(start_date)
            else:
                formatted_start = start_date

            if "-" in end_date and len(end_date.split("-")[0]) == 4:
                formatted_end = format_date(end_date)
            else:
                formatted_end = end_date

            data_payload = {
                "student_name": student_name,
                "student_email": recipient,
                "domain": domain,
                "duration": duration,
                "start_date": formatted_start,
                "end_date": formatted_end,
                "stipend": stipend,
            }
            if letter_type == "with_hours":
                weeks = DURATION_WEEKS.get(duration, 8)
                hours_per_week = int(data.get("hours_per_week") or 20)
                total_hours = weeks * hours_per_week
                data_payload["weeks"] = str(weeks)
                data_payload["hours_per_week"] = str(hours_per_week)
                data_payload["total_hours"] = str(total_hours)

            filename = generate_pdf(letter_type, data_payload)

        safe_fn = Path(filename).name
        pdf_path = GENERATED_DIR / safe_fn

        # --------------------------------------------------------
        # DUPLICATE CHECK & RESEND DETECTION
        # --------------------------------------------------------
        previous_record = get_previous_email_record(recipient)
        if previous_record and not send_again and str(previous_record.get("email_status") or "").lower() == "sent":
            return jsonify({
                "success": False,
                "duplicate": True,
                "skipped": True,
                "error": f"Email '{recipient}' already has a sent Offer Letter.",
                "previous_record": previous_record
            }), 409

        if not SENDER_EMAIL or not SENDER_PASSWORD:
            return jsonify({"success": False, "error": "Email is not configured on server."}), 500

        # Helper to construct history record
        def _build_bulk_history_data(status, sent_at=None, send_count=0, error_message=None):
            offer_letter_id = data.get("offer_letter_id") or Path(safe_fn).stem
            return {
                "student_name": student_name,
                "student_email": recipient,
                "phone_number": data.get("phone_number"),
                "college_name": data.get("college_name"),
                "internship_domain": domain,
                "internship_duration": duration,
                "start_date": start_date,
                "end_date": end_date,
                "offer_letter_type": letter_type,
                "email_status": status,
                "sent_at": sent_at,
                "offer_letter_id": offer_letter_id,
                "pdf_filename": safe_fn,
                "send_count": send_count,
                "error_message": error_message
            }

        # Prepare Email Message
        message = EmailMessage()
        message["From"] = f"Persevex LLP <{SENDER_EMAIL}>"
        message["To"] = recipient
        message["Subject"] = "Internship Acceptance Letter"

        message.set_content(
f"""Dear {student_name},

Warm greetings from Persevex LLP!

We are delighted to inform you that your Internship Acceptance Letter has been attached with this email. Please review the document carefully and feel free to reach out if you need any clarification.

We are truly excited to welcome you onboard at Persevex and look forward to your active contribution and learning journey with us. Your enthusiasm and dedication will play a key role in shaping meaningful experiences throughout this internship.

Kindly acknowledge the receipt of this email and confirm your acceptance at your earliest convenience.

Wishing you a wonderful start with us!

Warm regards,
Team Persevex"""
        )

        message.add_alternative(
f"""
<!DOCTYPE html>
<html>
<body style="font-family: Arial, Helvetica, sans-serif; font-size: 16px; line-height: 1.6; color: #333333;">

<p>Dear {student_name},</p>

<p>Warm greetings from Persevex LLP!</p>

<p>
We are delighted to inform you that your
<strong>Internship Acceptance Letter</strong>
has been attached with this email. Please review the document carefully
and feel free to reach out if you need any clarification.
</p>

<p>
We are truly excited to welcome you onboard at Persevex and look forward
to your active contribution and learning journey with us. Your enthusiasm
and dedication will play a key role in shaping meaningful experiences
throughout this internship.
</p>

<p>
Kindly acknowledge the receipt of this email and confirm your acceptance
at your earliest convenience.
</p>

<p>Wishing you a wonderful start with us!</p>

<p>
Warm regards,<br>
<strong>Team Persevex</strong>
</p>

</body>
</html>
""",
            subtype="html"
        )

        message.add_attachment(
            pdf_path.read_bytes(),
            maintype="application",
            subtype="pdf",
            filename=safe_fn
        )

        # Connect and Send
        try:
            with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30) as smtp:
                smtp.login(SENDER_EMAIL, SENDER_PASSWORD)
                smtp.send_message(message)
        except Exception as smtp_exc:
            err_msg = str(smtp_exc)
            try:
                if previous_record:
                    p_id = previous_record.get("id")
                    p_count = int(previous_record.get("send_count") or 0)
                    failed_data = _build_bulk_history_data(status="failed", send_count=p_count + 1, error_message=err_msg)
                    if supabase:
                        try:
                            supabase.table("email_history").update(failed_data).eq("id", p_id).execute()
                        except Exception:
                            pass
                else:
                    failed_data = _build_bulk_history_data(status="failed", send_count=1, error_message=err_msg)
                    repository.save_offer_letter_record(failed_data)
            except Exception:
                pass
            return jsonify({"success": False, "error": err_msg, "email_status": "failed", "student_email": recipient}), 500

        # Successful Send
        sent_time = datetime.now(timezone.utc).isoformat()
        is_resend = bool(previous_record and send_again)
        current_send_count = (int(previous_record.get("send_count") or 0) + 1) if is_resend else 1

        # Save to IMAP sent folder if possible
        try:
            with imaplib.IMAP4_SSL(IMAP_HOST, 993) as imap:
                imap.login(SENDER_EMAIL, SENDER_PASSWORD)
                imap.append(IMAP_SENT_FOLDER, "\\Seen", None, message.as_bytes())
        except Exception as imap_err:
            print("WARNING: Could not save bulk email to IMAP Sent folder:", repr(imap_err))

        # Save/Update history record
        if is_resend:
            try:
                p_id = previous_record.get("id")
                update_data = _build_bulk_history_data(status="sent", sent_at=sent_time, send_count=current_send_count, error_message=None)
                if supabase:
                    try:
                        supabase.table("email_history").update(update_data).eq("id", p_id).execute()
                    except Exception as exc:
                        print("SUPABASE BULK RESEND UPDATE ERROR:", repr(exc))
                try:
                    conn = sqlite3.connect(repository.SQLITE_DB_PATH)
                    cursor = conn.cursor()
                    cursor.execute("""
                        UPDATE email_history SET
                            student_name = ?, internship_domain = ?, internship_duration = ?,
                            start_date = ?, end_date = ?, offer_letter_type = ?,
                            email_status = 'sent', sent_at = ?, send_count = ?,
                            pdf_filename = ?, error_message = NULL
                        WHERE id = ? OR LOWER(TRIM(student_email)) = ?
                    """, (
                        student_name, domain, duration, start_date, end_date,
                        letter_type, sent_time, current_send_count, safe_fn,
                        p_id, recipient
                    ))
                    conn.commit()
                    conn.close()
                except Exception as db_exc:
                    print("SQLITE BULK RESEND UPDATE ERROR:", db_exc)
            except Exception as hist_err:
                print("BULK RESEND HISTORY ERROR:", repr(hist_err))
        else:
            new_history_data = _build_bulk_history_data(status="sent", sent_at=sent_time, send_count=1, error_message=None)
            repository.save_offer_letter_record(new_history_data)

        return jsonify({
            "success": True,
            "student_name": student_name,
            "student_email": recipient,
            "send_count": current_send_count,
            "resent": is_resend
        })

    except Exception as exc:
        print("BULK OFFER LETTER SEND ITEM ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.get("/api/offer-letter/bulk/sample-template")
def bulk_offer_letter_sample_template():
    """
    Download sample template for Bulk Offer Letters.
    """
    try:
        fmt = (request.args.get("format") or "csv").strip().lower()
        if fmt in ("excel", "xlsx"):
            excel_bytes = bulk_offer_letter_service.get_sample_excel_template()
            if excel_bytes:
                return Response(
                    excel_bytes,
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={
                        "Content-Disposition": 'attachment; filename="Persevex_Bulk_Offer_Letters_Sample.xlsx"',
                        "Cache-Control": "no-cache",
                    },
                )
        csv_bytes = bulk_offer_letter_service.get_sample_csv_template()
        return Response(
            csv_bytes,
            mimetype="text/csv",
            headers={
                "Content-Disposition": 'attachment; filename="Persevex_Bulk_Offer_Letters_Sample.csv"',
                "Cache-Control": "no-cache",
            },
        )
    except Exception as exc:
        print("BULK OFFER LETTER SAMPLE TEMPLATE ERROR:", repr(exc))
        return jsonify({"error": str(exc)}), 500


# ============================================================
# CAMPUS AMBASSADOR OFFER LETTER ROUTES
# ============================================================

@app.get("/campus-ambassador")
def campus_ambassador_menu_page():
    try:
        history_data = repository.get_unified_history(record_type="campus_ambassador", page=1, per_page=1)
        success_count = history_data.get("total_sent", 0)
        failed_count = history_data.get("total_failed", 0)
    except Exception as exc:
        print("CA HEADER STATUS COUNT ERROR:", repr(exc))
        success_count = 0
        failed_count = 0

    today_iso = datetime.now().strftime("%Y-%m-%d")
    today_formatted = datetime.now().strftime("%B %d, %Y")

    return render_template(
        "campus_ambassador.html",
        today_iso=today_iso,
        today_formatted=today_formatted,
        success_count=success_count,
        failed_count=failed_count,
    )


@app.get("/campus-ambassador/single")
def campus_ambassador_single_page():
    try:
        history_data = repository.get_unified_history(record_type="campus_ambassador", page=1, per_page=1)
        success_count = history_data.get("total_sent", 0)
        failed_count = history_data.get("total_failed", 0)
    except Exception as exc:
        print("CA HEADER STATUS COUNT ERROR:", repr(exc))
        success_count = 0
        failed_count = 0

    today_iso = datetime.now().strftime("%Y-%m-%d")
    today_formatted = datetime.now().strftime("%B %d, %Y")

    return render_template(
        "campus_ambassador.html",
        today_iso=today_iso,
        today_formatted=today_formatted,
        success_count=success_count,
        failed_count=failed_count,
    )


@app.get("/campus-ambassador/bulk")
def campus_ambassador_bulk_page():
    try:
        history_data = repository.get_unified_history(record_type="campus_ambassador", page=1, per_page=1)
        success_count = history_data.get("total_sent", 0)
        failed_count = history_data.get("total_failed", 0)
    except Exception as exc:
        print("CA HEADER STATUS COUNT ERROR:", repr(exc))
        success_count = 0
        failed_count = 0

    today_iso = datetime.now().strftime("%Y-%m-%d")
    today_formatted = datetime.now().strftime("%B %d, %Y")

    return render_template(
        "campus_ambassador_bulk.html",
        today_iso=today_iso,
        today_formatted=today_formatted,
        success_count=success_count,
        failed_count=failed_count,
    )


@app.get("/api/campus-ambassador/stats")
def campus_ambassador_stats():
    try:
        stats = repository.get_campus_ambassador_stats()
        return jsonify({
            "success": True,
            "success_count": stats.get("successful", 0),
            "failed_count": stats.get("failed", 0),
        })
    except Exception as exc:
        print("CA STATS API ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500



@app.post("/api/campus-ambassador/generate")
def campus_ambassador_generate():
    try:
        data = request.get_json(force=True) or {}
        student_name = (data.get("student_name") or "").strip()
        student_email = (data.get("student_email") or "").strip().lower()
        raw_date = data.get("date") or data.get("issue_date") or data.get("start_date") or ""

        if not student_name:
            return jsonify({"success": False, "error": "Candidate Name is required."}), 400

        if not student_email:
            return jsonify({"success": False, "error": "Candidate Email is required."}), 400

        date_formatted = format_ca_date(raw_date)

        data_payload = {
            "student_name": student_name,
            "student_email": student_email,
            "date": date_formatted,
            "domain": "Campus Ambassador",
            "offer_letter_type": "campus_ambassador",
        }

        filename = generate_ca_pdf(data_payload)

        return jsonify({
            "success": True,
            "filename": filename,
            "url": f"/generated/{filename}",
            "student_name": student_name,
            "student_email": student_email,
            "date": date_formatted,
        })
    except Exception as exc:
        print("CAMPUS AMBASSADOR GENERATION ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


def get_ca_email_html(student_name):
    """
    Generate clean, professional business email HTML for Campus Ambassador Offer Letters.
    Left-aligned, standard typography, natural email width, no webpage/card/box borders.
    """
    return f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8">
</head>
<body style="font-family: Arial, Helvetica, sans-serif; font-size: 14px; line-height: 1.6; color: #1e293b; margin: 0; padding: 10px 0; background-color: #ffffff;">
  <div style="font-family: Arial, Helvetica, sans-serif; font-size: 14px; line-height: 1.6; color: #1e293b;">
    <p style="margin-top: 0; margin-bottom: 16px; color: #0f172a; font-weight: bold; font-size: 15px;">Dear {student_name},</p>
    <p style="margin-top: 0; margin-bottom: 16px;">Greetings from Persevex!</p>
    <p style="margin-top: 0; margin-bottom: 16px;">We are excited to officially welcome you as a Campus Ambassador at Persevex. Please find attached your appointment letter, which outlines your key responsibilities, benefits, and the impact you can make as part of our team.</p>
    <p style="margin-top: 0; margin-bottom: 16px;">As a Campus Ambassador, you will play a vital role in building brand awareness, promoting our programs, and fostering student engagement at your institution. Your energy and initiative will be instrumental in expanding Persevex’s mission to empower learners across campuses.</p>
    <p style="margin-top: 0; margin-bottom: 16px;">If you have any questions or need further clarification, feel free to reach out to us at 📧 <a href="mailto:support@persevex.com" style="color: #2563eb; text-decoration: none;">support@persevex.com</a>.</p>
    <p style="margin-top: 0; margin-bottom: 16px;">We look forward to seeing your contributions and success in this role.</p>
    <p style="margin-top: 0; margin-bottom: 24px;">📣 Feel free to share this exciting opportunity on LinkedIn by posting about your new role, tagging @Persevex and using hashtags such as #Persevex #CampusAmbassador #Leadership #StudentOpportunity #EmpoweringLearners.</p>
    <p style="margin-top: 0; margin-bottom: 4px;">Best regards,</p>
    <p style="margin-top: 0; margin-bottom: 0;">
      <strong>Team Persevex</strong><br>
      📧 <a href="mailto:support@persevex.com" style="color: #2563eb; text-decoration: none;">support@persevex.com</a><br>
      🌐 <a href="https://www.persevex.com" style="color: #2563eb; text-decoration: none;">www.persevex.com</a>
    </p>
  </div>
</body>
</html>"""


@app.post("/api/campus-ambassador/send-email")
def campus_ambassador_send_email():
    data = request.get_json(force=True) or {}
    filename = data.get("filename") or ""
    recipient = (data.get("student_email") or "").strip().lower()
    student_name = (data.get("student_name") or "Candidate").strip()
    raw_date = data.get("date") or data.get("issue_date") or data.get("start_date") or ""
    date_formatted = format_ca_date(raw_date)
    send_again = data.get("send_again") is True or data.get("force_resend") is True

    if not filename:
        return jsonify({"success": False, "error": "Generated PDF filename is required."}), 400

    if not recipient:
        return jsonify({"success": False, "error": "Candidate email is required."}), 400

    filename = Path(filename).name
    pdf_path = GENERATED_DIR / filename

    if not pdf_path.exists():
        return jsonify({
            "success": False,
            "error": "Generated PDF not found on server. Please generate the letter again."
        }), 404

    # --------------------------------------------------------
    # DUPLICATE / PREVIOUSLY SENT CHECK (public.campus_ambassador_history ONLY)
    # --------------------------------------------------------
    previous_sent_record = repository.get_previous_sent_campus_ambassador_by_email(recipient)

    if previous_sent_record and not send_again:
        return jsonify({
            "success": False,
            "duplicate": True,
            "previously_sent": True,
            "message": "This email has already received a Campus Ambassador Offer Letter. Do you want to send it again?",
            "previous_record": previous_sent_record
        }), 409

    is_resend = bool(previous_sent_record and send_again)
    target_record_id = previous_sent_record.get("id") if is_resend else None
    current_send_count = ((previous_sent_record.get("send_count") or 1) + 1) if is_resend else 1

    safe_fn = Path(filename).name
    doc_id = Path(safe_fn).stem

    def build_ca_history_data(status, sent_at=None, send_count=current_send_count, error_message=None):
        payload = {
            "student_name": student_name,
            "student_email": recipient,
            "internship_domain": "Campus Ambassador",
            "internship_duration": "Tenure",
            "start_date": date_formatted,
            "end_date": None,
            "offer_letter_type": "campus_ambassador",
            "email_status": status,
            "sent_at": sent_at,
            "offer_letter_id": doc_id,
            "pdf_filename": safe_fn,
            "send_count": send_count,
            "error_message": error_message,
        }
        if is_resend and target_record_id:
            payload["id"] = target_record_id
            payload["send_again"] = True
        return payload

    try:
        message = EmailMessage()
        message["Subject"] = "Appointment Letter – Campus Ambassador at Persevex"
        message["From"] = f"Persevex LLP <{SENDER_EMAIL}>"
        message["To"] = recipient

        html_content = get_ca_email_html(student_name)

        plain_text = (
            f"Dear {student_name},\n\n"
            "Greetings from Persevex!\n\n"
            "We are excited to officially welcome you as a Campus Ambassador at Persevex. Please find attached your appointment letter, which outlines your key responsibilities, benefits, and the impact you can make as part of our team.\n\n"
            "As a Campus Ambassador, you will play a vital role in building brand awareness, promoting our programs, and fostering student engagement at your institution. Your energy and initiative will be instrumental in expanding Persevex’s mission to empower learners across campuses.\n\n"
            "If you have any questions or need further clarification, feel free to reach out to us at 📧 support@persevex.com.\n\n"
            "We look forward to seeing your contributions and success in this role.\n\n"
            "📣 Feel free to share this exciting opportunity on LinkedIn by posting about your new role, tagging @Persevex and using hashtags such as #Persevex #CampusAmbassador #Leadership #StudentOpportunity #EmpoweringLearners.\n\n"
            "Best regards,\n"
            "Team Persevex\n"
            "📧 support@persevex.com\n"
            "🌐 www.persevex.com"
        )

        message.set_content(plain_text)
        message.add_alternative(html_content, subtype="html")

        message.add_attachment(
            pdf_path.read_bytes(),
            maintype="application",
            subtype="pdf",
            filename=filename,
        )
    except Exception as exc:
        print("CA EMAIL PREPARATION ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc), "email_status": "failed"}), 500

    try:
        print("========================================")
        print("CONNECTING TO SMTP (CA LETTER)")
        print("FROM:", SENDER_EMAIL)
        print("TO:", recipient)
        print("========================================")

        with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30) as smtp:
            smtp.login(SENDER_EMAIL, SENDER_PASSWORD)
            smtp.send_message(message)

        print("SMTP CA EMAIL SENT SUCCESSFULLY")
    except Exception as exc:
        error_msg = str(exc)
        print("SMTP CA EMAIL FAILED:", repr(exc))
        try:
            failed_record = build_ca_history_data(status="failed", error_message=error_msg)
            repository.save_campus_ambassador_record(failed_record)
        except Exception as save_err:
            print("CA FAILED HISTORY SAVE ERROR:", repr(save_err))
        return jsonify({"success": False, "error": error_msg, "email_status": "failed"}), 500

    sent_time = datetime.now(timezone.utc).isoformat()

    try:
        with imaplib.IMAP4_SSL(IMAP_HOST, IMAP_PORT) as imap:
            imap.login(SENDER_EMAIL, SENDER_PASSWORD)
            imap.append(IMAP_SENT_FOLDER, "\\Seen", None, message.as_bytes())
    except Exception as imap_err:
        print("WARNING: Could not save CA email to IMAP Sent folder:", repr(imap_err))

    try:
        success_record = build_ca_history_data(status="sent", sent_at=sent_time, send_count=current_send_count)
        repository.save_campus_ambassador_record(success_record, existing_id=target_record_id if is_resend else None)
    except Exception as save_err:
        print("CA SUCCESS HISTORY SAVE ERROR:", repr(save_err))

    return jsonify({
        "success": True,
        "message": "Campus Ambassador Offer Letter sent successfully!",
        "send_count": current_send_count,
        "resent": is_resend
    })


@app.post("/api/campus-ambassador/parse-upload")
def campus_ambassador_parse_upload():
    """Parse uploaded CSV or Excel file for Bulk Campus Ambassador offer letters."""
    if "file" not in request.files:
        return jsonify({"success": False, "error": "No file uploaded."}), 400
    file = request.files["file"]
    if not file or not file.filename:
        return jsonify({"success": False, "error": "No file selected."}), 400

    file_bytes = file.read()
    result = parse_ca_file(file_bytes, file.filename)
    if not result.get("success"):
        return jsonify(result), 400
    return jsonify(result)


@app.post("/api/campus-ambassador/bulk-generate-item")
def campus_ambassador_bulk_generate_item():
    """Generate a single CA Offer Letter PDF without dispatching email."""
    try:
        data = request.get_json(force=True) or {}
        student_name = (data.get("student_name") or "").strip()
        student_email = (data.get("student_email") or "").strip().lower()
        raw_date = data.get("date") or data.get("issue_date") or data.get("start_date") or ""

        if not student_name:
            return jsonify({"success": False, "error": "Candidate Name is required."}), 400
        if not student_email:
            return jsonify({"success": False, "error": "Candidate Email is required."}), 400

        date_formatted = format_ca_date(raw_date)
        data_payload = {
            "student_name": student_name,
            "student_email": student_email,
            "date": date_formatted,
            "domain": "Campus Ambassador",
            "offer_letter_type": "campus_ambassador",
        }
        filename = generate_ca_pdf(data_payload)
        return jsonify({
            "success": True,
            "filename": filename,
            "student_name": student_name,
            "student_email": student_email,
            "date": date_formatted
        })
    except Exception as exc:
        print("BULK CA GENERATE ITEM ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.post("/api/campus-ambassador/bulk-check-duplicates")
def campus_ambassador_bulk_check_duplicates():
    """
    Batch duplicate check against public.campus_ambassador_history.
    Considers ONLY records where email_status = 'sent'.
    Normalizes candidate emails and database emails using email.strip().lower().
    Authoritative source: Supabase (falls back to local SQLite on exception/unconfigured).
    Returns list of duplicate email strings and mapping of details.
    """
    try:
        data = request.get_json(force=True) or {}
        records = data.get("records") or data.get("emails") or []
        if not records:
            return jsonify({
                "success": True,
                "duplicate_count": 0,
                "duplicates": [],
                "duplicate_map": {}
            })

        duplicate_map = repository.find_existing_campus_ambassador_batch(records)
        duplicate_emails = list(duplicate_map.keys())

        return jsonify({
            "success": True,
            "duplicate_count": len(duplicate_emails),
            "duplicates": duplicate_emails,
            "duplicate_map": duplicate_map
        })
    except Exception as exc:
        print("BULK CA CHECK DUPLICATES ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.post("/api/campus-ambassador/bulk-send-item")
def campus_ambassador_bulk_send_item():
    """Dispatch email for a generated CA Offer Letter PDF."""
    student_name = ""
    student_email = ""
    try:
        data = request.get_json(force=True) or {}
        filename = (data.get("filename") or "").strip()
        student_name = (data.get("student_name") or "").strip()
        student_email = (data.get("student_email") or "").strip().lower()
        raw_date = data.get("date") or data.get("issue_date") or data.get("start_date") or ""
        send_again = bool(data.get("send_again", False))
        prev_count_param = data.get("previous_send_count")

        if not filename:
            # Fallback generate PDF if not provided
            date_formatted = format_ca_date(raw_date)
            data_payload = {
                "student_name": student_name,
                "student_email": student_email,
                "date": date_formatted,
                "domain": "Campus Ambassador",
                "offer_letter_type": "campus_ambassador",
            }
            filename = generate_ca_pdf(data_payload)

        safe_fn = Path(filename).name
        pdf_path = GENERATED_DIR / safe_fn
        if not pdf_path.exists():
            return jsonify({
                "success": False,
                "error": "Generated PDF not found on server. Please regenerate.",
                "student_name": student_name,
                "student_email": student_email
            }), 404

        date_formatted = format_ca_date(raw_date)
        doc_id = Path(safe_fn).stem

        # --------------------------------------------------------
        # BACKEND DUPLICATE GUARD (public.campus_ambassador_history)
        # --------------------------------------------------------
        previous_sent_record = None
        if not send_again:
            previous_sent_record = repository.get_previous_sent_campus_ambassador_by_email(student_email)
            if previous_sent_record and str(previous_sent_record.get("email_status") or "").strip().lower() == "sent":
                return jsonify({
                    "success": False,
                    "duplicate": True,
                    "skipped": True,
                    "error": f"Email '{student_email}' has already received a Campus Ambassador Offer Letter.",
                    "student_name": student_name,
                    "student_email": student_email
                }), 409
        else:
            previous_sent_record = repository.get_previous_sent_campus_ambassador_by_email(student_email)

        is_resend = bool(send_again and previous_sent_record and str(previous_sent_record.get("email_status") or "").strip().lower() == "sent")
        target_record_id = (data.get("id") or data.get("existing_id") or (previous_sent_record.get("id") if previous_sent_record else None)) if is_resend else None

        if is_resend:
            if prev_count_param is not None and int(prev_count_param) >= 1:
                current_send_count = int(prev_count_param) + 1
            elif previous_sent_record:
                current_send_count = (int(previous_sent_record.get("send_count") or 1)) + 1
            else:
                current_send_count = 2
        else:
            current_send_count = 1

        def build_ca_history_data(status, sent_at=None, send_count=current_send_count, error_message=None):
            payload = {
                "student_name": student_name,
                "student_email": student_email,
                "internship_domain": "Campus Ambassador",
                "internship_duration": "Tenure",
                "start_date": date_formatted,
                "end_date": None,
                "offer_letter_type": "campus_ambassador",
                "email_status": status,
                "sent_at": sent_at,
                "offer_letter_id": doc_id,
                "pdf_filename": safe_fn,
                "send_count": send_count,
                "error_message": error_message,
            }
            if is_resend and target_record_id:
                payload["id"] = target_record_id
                payload["send_again"] = True
            return payload

        message = EmailMessage()
        message["Subject"] = "Appointment Letter – Campus Ambassador at Persevex"
        message["From"] = f"Persevex LLP <{SENDER_EMAIL}>"
        message["To"] = student_email

        html_content = get_ca_email_html(student_name)
        plain_text = (
            f"Dear {student_name},\n\n"
            "Greetings from Persevex!\n\n"
            "We are excited to officially welcome you as a Campus Ambassador at Persevex. Please find attached your appointment letter, which outlines your key responsibilities, benefits, and the impact you can make as part of our team.\n\n"
            "As a Campus Ambassador, you will play a vital role in building brand awareness, promoting our programs, and fostering student engagement at your institution. Your energy and initiative will be instrumental in expanding Persevex’s mission to empower learners across campuses.\n\n"
            "If you have any questions or need further clarification, feel free to reach out to us at 📧 support@persevex.com.\n\n"
            "We look forward to seeing your contributions and success in this role.\n\n"
            "📣 Feel free to share this exciting opportunity on LinkedIn by posting about your new role, tagging @Persevex and using hashtags such as #Persevex #CampusAmbassador #Leadership #StudentOpportunity #EmpoweringLearners.\n\n"
            "Best regards,\n"
            "Team Persevex\n"
            "📧 support@persevex.com\n"
            "🌐 www.persevex.com"
        )
        message.set_content(plain_text)
        message.add_alternative(html_content, subtype="html")
        message.add_attachment(
            pdf_path.read_bytes(),
            maintype="application",
            subtype="pdf",
            filename=safe_fn,
        )

        try:
            with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30) as smtp:
                smtp.login(SENDER_EMAIL, SENDER_PASSWORD)
                smtp.send_message(message)
        except Exception as smtp_exc:
            error_msg = str(smtp_exc)
            print("BULK CA SMTP EMAIL FAILED:", repr(smtp_exc))
            try:
                failed_record = build_ca_history_data(status="failed", error_message=error_msg)
                repository.save_campus_ambassador_record(failed_record)
            except Exception as save_err:
                print("BULK CA FAILED HISTORY SAVE ERROR:", repr(save_err))
            return jsonify({
                "success": False,
                "student_name": student_name,
                "student_email": student_email,
                "error": error_msg,
                "email_status": "failed"
            }), 500

        sent_time = datetime.now(timezone.utc).isoformat()
        try:
            with imaplib.IMAP4_SSL(IMAP_HOST, IMAP_PORT) as imap:
                imap.login(SENDER_EMAIL, SENDER_PASSWORD)
                imap.append(IMAP_SENT_FOLDER, "\\Seen", None, message.as_bytes())
        except Exception:
            pass

        try:
            success_record = build_ca_history_data(status="sent", sent_at=sent_time, send_count=current_send_count)
            repository.save_campus_ambassador_record(success_record, existing_id=target_record_id if is_resend else None)
        except Exception as save_err:
            print("BULK CA SUCCESS HISTORY SAVE ERROR:", repr(save_err))

        return jsonify({
            "success": True,
            "student_name": student_name,
            "student_email": student_email,
            "filename": safe_fn,
            "send_count": current_send_count,
            "resent": is_resend,
            "message": "Campus Ambassador Offer Letter sent successfully!"
        })

    except Exception as exc:
        print("BULK CA ITEM SEND ERROR:", repr(exc))
        return jsonify({
            "success": False,
            "student_name": student_name,
            "student_email": student_email,
            "error": str(exc)
        }), 500


@app.post("/api/campus-ambassador/bulk-process-item")
def campus_ambassador_bulk_process_item():
    student_name = ""
    student_email = ""
    try:
        data = request.get_json(force=True) or {}
        student_name = (data.get("student_name") or "").strip()
        student_email = (data.get("student_email") or "").strip().lower()
        raw_date = data.get("date") or data.get("issue_date") or data.get("start_date") or ""

        if not student_name:
            return jsonify({
                "success": False,
                "error": "Candidate Name is required.",
                "student_name": student_name,
                "student_email": student_email
            }), 400

        if not student_email:
            return jsonify({
                "success": False,
                "error": "Candidate Email is required.",
                "student_name": student_name,
                "student_email": student_email
            }), 400

        date_formatted = format_ca_date(raw_date)

        # 1. Generate PDF
        data_payload = {
            "student_name": student_name,
            "student_email": student_email,
            "date": date_formatted,
            "domain": "Campus Ambassador",
            "offer_letter_type": "campus_ambassador",
        }
        filename = generate_ca_pdf(data_payload)
        safe_fn = Path(filename).name
        doc_id = Path(safe_fn).stem
        pdf_path = GENERATED_DIR / safe_fn

        def build_ca_history_data(status, sent_at=None, send_count=1, error_message=None):
            return {
                "student_name": student_name,
                "student_email": student_email,
                "internship_domain": "Campus Ambassador",
                "internship_duration": "Tenure",
                "start_date": date_formatted,
                "end_date": None,
                "offer_letter_type": "campus_ambassador",
                "email_status": status,
                "sent_at": sent_at,
                "offer_letter_id": doc_id,
                "pdf_filename": safe_fn,
                "send_count": send_count,
                "error_message": error_message,
            }

        # 2. Prepare Email
        message = EmailMessage()
        message["Subject"] = "Appointment Letter – Campus Ambassador at Persevex"
        message["From"] = f"Persevex LLP <{SENDER_EMAIL}>"
        message["To"] = student_email

        html_content = get_ca_email_html(student_name)

        plain_text = (
            f"Dear {student_name},\n\n"
            "Greetings from Persevex!\n\n"
            "We are excited to officially welcome you as a Campus Ambassador at Persevex. Please find attached your appointment letter, which outlines your key responsibilities, benefits, and the impact you can make as part of our team.\n\n"
            "As a Campus Ambassador, you will play a vital role in building brand awareness, promoting our programs, and fostering student engagement at your institution. Your energy and initiative will be instrumental in expanding Persevex’s mission to empower learners across campuses.\n\n"
            "If you have any questions or need further clarification, feel free to reach out to us at 📧 support@persevex.com.\n\n"
            "We look forward to seeing your contributions and success in this role.\n\n"
            "📣 Feel free to share this exciting opportunity on LinkedIn by posting about your new role, tagging @Persevex and using hashtags such as #Persevex #CampusAmbassador #Leadership #StudentOpportunity #EmpoweringLearners.\n\n"
            "Best regards,\n"
            "Team Persevex\n"
            "📧 support@persevex.com\n"
            "🌐 www.persevex.com"
        )

        message.set_content(plain_text)
        message.add_alternative(html_content, subtype="html")

        message.add_attachment(
            pdf_path.read_bytes(),
            maintype="application",
            subtype="pdf",
            filename=safe_fn,
        )

        # 3. Send Email via SMTP
        try:
            with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30) as smtp:
                smtp.login(SENDER_EMAIL, SENDER_PASSWORD)
                smtp.send_message(message)
        except Exception as smtp_exc:
            error_msg = str(smtp_exc)
            print("BULK CA SMTP EMAIL FAILED:", repr(smtp_exc))
            try:
                failed_record = build_ca_history_data(status="failed", error_message=error_msg)
                repository.save_campus_ambassador_record(failed_record)
            except Exception as save_err:
                print("BULK CA FAILED HISTORY SAVE ERROR:", repr(save_err))
            return jsonify({
                "success": False,
                "student_name": student_name,
                "student_email": student_email,
                "error": error_msg,
                "email_status": "failed"
            }), 500

        sent_time = datetime.now(timezone.utc).isoformat()

        # 4. IMAP backup
        try:
            with imaplib.IMAP4_SSL(IMAP_HOST, IMAP_PORT) as imap:
                imap.login(SENDER_EMAIL, SENDER_PASSWORD)
                imap.append(IMAP_SENT_FOLDER, "\\Seen", None, message.as_bytes())
        except Exception as imap_err:
            print("WARNING: Could not save bulk CA email to IMAP Sent folder:", repr(imap_err))

        # 5. Save success history
        try:
            success_record = build_ca_history_data(status="sent", sent_at=sent_time, send_count=1)
            repository.save_campus_ambassador_record(success_record)
        except Exception as save_err:
            print("BULK CA SUCCESS HISTORY SAVE ERROR:", repr(save_err))

        return jsonify({
            "success": True,
            "student_name": student_name,
            "student_email": student_email,
            "filename": safe_fn,
            "message": "Campus Ambassador Offer Letter generated and sent successfully!"
        })

    except Exception as exc:
        print("BULK CA ITEM PROCESSING ERROR:", repr(exc))
        return jsonify({
            "success": False,
            "student_name": student_name,
            "student_email": student_email,
            "error": str(exc)
        }), 500



# ============================================================
# UNIFIED HISTORY ROUTES
# ============================================================

@app.get("/history")
def history():
    try:
        search = (request.args.get("search", "") or "").strip()
        record_type = (request.args.get("record_type", "") or request.args.get("type", "") or "all").strip().lower()
        status_filter = (request.args.get("status", "") or "").strip().lower()
        month_filter = (request.args.get("month", "") or "").strip()
        date_filter = (request.args.get("date_filter", "") or request.args.get("date", "") or "all").strip().lower()
        date_from = (request.args.get("date_from", "") or "").strip()
        date_to = (request.args.get("date_to", "") or "").strip()

        try:
            per_page = int(request.args.get("per_page", 25))
            if per_page not in [10, 25, 50, 100, 200]:
                per_page = 25
        except (TypeError, ValueError):
            per_page = 25

        try:
            page = int(request.args.get("page", 1))
        except (TypeError, ValueError):
            page = 1
        if page < 1:
            page = 1

        history_data = repository.get_unified_history(
            search=search,
            record_type=record_type,
            status_filter=status_filter,
            month_filter=month_filter,
            date_filter=date_filter,
            date_from=date_from,
            date_to=date_to,
            page=page,
            per_page=per_page,
        )

        is_json = (
            request.args.get("format") == "json"
            or request.headers.get("Accept") == "application/json"
            or request.headers.get("X-Requested-With") == "XMLHttpRequest"
        )

        if is_json:
            return jsonify({
                "success": True,
                "records": history_data["records"],
                "total_records": history_data["total_records"],
                "total_offer_letters": history_data["total_offer_letters"],
                "total_ca_letters": history_data.get("total_ca_letters", 0),
                "total_certificates": history_data["total_certificates"],
                "total_ca_certificates": history_data.get("total_ca_certificates", 0),
                "total_sent": history_data["total_sent"],
                "total_failed": history_data["total_failed"],
                "total_pending": history_data["total_pending"],
                "page": history_data["page"],
                "per_page": history_data["per_page"],
                "total_pages": history_data["total_pages"],
                "available_months": history_data["available_months"],
            })

        return render_template(
            "history.html",
            history_records=history_data["records"],
            search=search,
            record_type=record_type,
            status_filter=status_filter,
            month_filter=month_filter,
            date_filter=date_filter,
            date_from=date_from,
            date_to=date_to,
            page=history_data["page"],
            total_pages=history_data["total_pages"],
            total_records=history_data["total_records"],
            total_offer_letters=history_data["total_offer_letters"],
            total_ca_letters=history_data.get("total_ca_letters", 0),
            total_certificates=history_data["total_certificates"],
            total_ca_certificates=history_data.get("total_ca_certificates", 0),
            total_sent=history_data["total_sent"],
            total_failed=history_data["total_failed"],
            total_pending=history_data["total_pending"],
            per_page=history_data["per_page"],
            available_months=history_data["available_months"],
            error=None,
        )

    except Exception as exc:
        print("UNIFIED HISTORY ERROR:", repr(exc))
        is_json = (
            request.args.get("format") == "json"
            or request.headers.get("Accept") == "application/json"
            or request.headers.get("X-Requested-With") == "XMLHttpRequest"
        )
        if is_json:
            return jsonify({
                "success": False,
                "error": str(exc),
                "records": [],
                "total_records": 0,
                "total_offer_letters": 0,
                "total_ca_letters": 0,
                "total_certificates": 0,
                "total_sent": 0,
                "total_failed": 0,
                "page": 1,
                "total_pages": 1,
                "per_page": 25,
                "available_months": [],
            }), 500

        return render_template(
            "history.html",
            history_records=[],
            search="",
            record_type="all",
            status_filter="",
            month_filter="",
            date_filter="all",
            date_from="",
            date_to="",
            page=1,
            total_pages=1,
            total_records=0,
            total_offer_letters=0,
            total_ca_letters=0,
            total_certificates=0,
            total_ca_certificates=0,
            total_sent=0,
            total_failed=0,
            total_pending=0,
            per_page=25,
            available_months=[],
            error=str(exc),
        )


@app.get("/history/export")
def export_history():
    try:
        record_type = (request.args.get("record_type") or request.args.get("type") or "all").strip().lower()
        from_date = (request.args.get("date_from") or request.args.get("from_date") or "").strip()
        to_date = (request.args.get("date_to") or request.args.get("to_date") or "").strip()
        status_filter = (request.args.get("status") or "").strip().lower()
        export_format = (request.args.get("format") or "csv").strip().lower()

        if from_date and to_date and from_date > to_date:
            return jsonify({"success": False, "error": "From Date cannot be later than To Date."}), 400

        history_data = repository.get_unified_history(
            search="",
            record_type=record_type,
            status_filter=status_filter,
            date_from=from_date,
            date_to=to_date,
            page=1,
            per_page=100000,
        )

        records = history_data.get("all_filtered_records") or history_data.get("records") or []

        filename_prefix = "Persevex_Unified_History"
        if record_type in ("offer_letter", "offer", "offer_letters"):
            filename_prefix = "Persevex_Offer_Letter_History"
        elif record_type in ("ca_letter", "ca", "campus_ambassador", "ca_offer_letter", "ca_offer_letters"):
            filename_prefix = "Persevex_CA_Offer_Letter_History"
        elif record_type in ("certificate", "cert", "certificates"):
            filename_prefix = "Persevex_Certificate_History"
        elif record_type in ("ca_certificate", "ca_certificates", "ca_cert"):
            filename_prefix = "Persevex_CA_Certificate_History"

        date_suffix = datetime.now().strftime("%Y-%m-%d")
        if from_date and to_date:
            date_suffix = f"{from_date}_to_{to_date}"
        elif from_date:
            date_suffix = f"from_{from_date}"

        # 1. EXCEL (.xlsx) EXPORT
        if export_format in ("excel", "xlsx"):
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "History Audit"

                headers = [
                    "Type", "Document ID", "Student Name", "Student Email",
                    "Domain", "Duration / Track", "Start Date", "End Date",
                    "Issued Date", "Email Status", "Sent Date & Time", "Send Count", "PDF Filename"
                ]
                ws.append(headers)

                header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                for r in records:
                    raw_ts = str(r.get("sent_at") or r.get("created_at") or "")
                    date_display = raw_ts[:19].replace("T", " ") if raw_ts else "-"

                    status_val = str(r.get("email_status") or "").upper()
                    if status_val == "SENT":
                        status_display = "SENT"
                    elif status_val == "FAILED":
                        status_display = "FAILED"
                    else:
                        status_display = "PENDING"

                    ws.append([
                        r.get("type_label") or ("Offer Letter" if r.get("record_type") == "offer_letter" else "Certificate"),
                        r.get("document_id") or "-",
                        r.get("student_name") or "-",
                        r.get("student_email") or "-",
                        r.get("domain") or "-",
                        r.get("duration") or r.get("letter_type") or "-",
                        r.get("start_date") or "-",
                        r.get("end_date") or "-",
                        r.get("issued_date") or "-",
                        status_display,
                        date_display,
                        r.get("send_count") or 0,
                        r.get("pdf_filename") or "-",
                    ])

                # Adjust column widths
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

                excel_buffer = io.BytesIO()
                wb.save(excel_buffer)
                excel_bytes = excel_buffer.getvalue()
                excel_buffer.close()

                export_filename = f"{filename_prefix}_{date_suffix}.xlsx"
                return Response(
                    excel_bytes,
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={
                        "Content-Disposition": f'attachment; filename="{export_filename}"',
                        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    }
                )
            except Exception as excel_err:
                print("EXCEL EXPORT FALLBACK TO CSV:", repr(excel_err))
                # Fall through to CSV

        # 2. CSV EXPORT
        output = io.StringIO()
        output.write("\ufeff")  # UTF-8 BOM
        writer = csv.writer(output)

        writer.writerow([
            "Type", "Document ID", "Student Name", "Student Email",
            "Domain", "Duration / Track", "Start Date", "End Date",
            "Issued Date", "Email Status", "Sent Date & Time", "Send Count", "PDF Filename"
        ])

        for r in records:
            raw_ts = str(r.get("sent_at") or r.get("created_at") or "")
            date_display = raw_ts[:19].replace("T", " ") if raw_ts else "-"
            status_val = str(r.get("email_status") or "").upper()
            status_display = "SENT" if status_val == "SENT" else ("FAILED" if status_val == "FAILED" else "PENDING")

            writer.writerow([
                r.get("type_label") or ("Offer Letter" if r.get("record_type") == "offer_letter" else "Certificate"),
                r.get("document_id") or "-",
                r.get("student_name") or "-",
                r.get("student_email") or "-",
                r.get("domain") or "-",
                r.get("duration") or r.get("letter_type") or "-",
                r.get("start_date") or "-",
                r.get("end_date") or "-",
                r.get("issued_date") or "-",
                status_display,
                date_display,
                r.get("send_count") or 0,
                r.get("pdf_filename") or "-",
            ])

        export_filename = f"{filename_prefix}_{date_suffix}.csv"
        csv_content = output.getvalue()
        output.close()

        return Response(
            csv_content,
            mimetype="text/csv",
            headers={
                "Content-Disposition": f'attachment; filename="{export_filename}"',
                "Content-Type": "text/csv; charset=utf-8",
            }
        )

    except Exception as exc:
        print("EXPORT HISTORY ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.delete("/api/history/<string:record_type>/<path:record_id>")
def delete_unified_history_record(record_type, record_id):
    try:
        rec_type = str(record_type or "").strip().lower()
        if rec_type in ("ca_letter", "ca", "campus_ambassador", "ca_offer_letter", "ca_offer_letters"):
            deleted = repository.delete_campus_ambassador_record(record_id)
        elif rec_type in ("ca_certificate", "ca_certificates", "ca_cert"):
            deleted = repository.delete_ca_certificate_record(record_id)
        elif rec_type in ("offer_letter", "offer", "offer_letters", "email_history"):
            deleted = repository.delete_offer_letter_record(record_id)
        else:
            deleted = repository.delete_certificate_record(record_id)

        if not deleted:
            return jsonify({
                "success": False,
                "error": "Record not found or already deleted."
            }), 404

        return jsonify({
            "success": True,
            "message": "Audit record deleted successfully."
        })
    except Exception as exc:
        print("DELETE UNIFIED RECORD ERROR:", repr(exc))
        return jsonify({"success": False, "error": "Failed to delete record. Please check server logs."}), 500


@app.delete("/history/<int:record_id>")
def delete_history_record(record_id):
    try:
        deleted = repository.delete_offer_letter_record(record_id)
        if not deleted:
            return jsonify({"success": False, "error": "Record not found or already deleted."}), 404
        return jsonify({
            "success": True,
            "message": "Offer letter history record deleted successfully."
        })
    except Exception as exc:
        print("DELETE HISTORY ERROR:", repr(exc))
        return jsonify({"success": False, "error": "Failed to delete record. Please check server logs."}), 500


@app.post("/api/history/bulk-delete")
def bulk_delete_history_records():
    try:
        data = request.get_json(force=True) or {}
        items = data.get("items") or []
        if not items:
            return jsonify({"success": False, "error": "No records selected for deletion."}), 400

        result = repository.bulk_delete_records(items)
        return jsonify({
            "success": True,
            "deleted_count": result.get("total_deleted", 0),
            "offer_letters_deleted": result.get("offer_letters_deleted", 0),
            "ca_letters_deleted": result.get("ca_letters_deleted", 0),
            "certificates_deleted": result.get("certificates_deleted", 0),
            "ca_certificates_deleted": result.get("ca_certificates_deleted", 0),
            "message": f"{result.get('total_deleted', 0)} records deleted successfully."
        })
    except Exception as exc:
        print("BULK DELETE API ERROR:", repr(exc))
        return jsonify({"success": False, "error": "Failed to delete records. Please check server logs."}), 500


@app.post("/api/history/delete-all")
def delete_all_history_endpoint():
    try:
        data = request.get_json(force=True) or {}
        confirm_phrase = str(data.get("confirm_phrase") or "").strip().upper()
        record_type = str(data.get("record_type") or "all").strip().lower()

        if confirm_phrase != "DELETE":
            return jsonify({
                "success": False,
                "error": "Safety check failed. Please type DELETE to confirm."
            }), 400

        result = repository.delete_all_history_records(record_type)
        return jsonify({
            "success": True,
            "deleted_count": result.get("total_deleted", 0),
            "offer_letters_deleted": result.get("offer_letters_deleted", 0),
            "ca_letters_deleted": result.get("ca_letters_deleted", 0),
            "certificates_deleted": result.get("certificates_deleted", 0),
            "ca_certificates_deleted": result.get("ca_certificates_deleted", 0),
            "message": "All history records deleted successfully."
        })
    except Exception as exc:
        print("DELETE ALL API ERROR:", repr(exc))
        return jsonify({"success": False, "error": "Failed to purge history records. Please check server logs."}), 500


# ============================================================
# CERTIFICATE GENERATOR & BULK ROUTES
# ============================================================

@app.get("/ca-certificate")
@app.get("/ca-certificate/single")
def ca_certificate_page():
    stats = repository.get_unified_history(record_type="ca_certificate", page=1, per_page=1)
    return render_template(
        "ca_certificate.html",
        today_iso=datetime.now(APP_LOCAL_TIMEZONE).strftime("%Y-%m-%d"),
        today_formatted=datetime.now(APP_LOCAL_TIMEZONE).strftime("%B %d, %Y"),
        success_count=stats.get("total_sent", 0),
        failed_count=stats.get("total_failed", 0),
    )


@app.get("/ca-certificate/bulk")
def ca_certificate_bulk_page():
    stats = repository.get_unified_history(record_type="ca_certificate", page=1, per_page=1)
    return render_template(
        "ca_certificate_bulk.html",
        today_iso=datetime.now(APP_LOCAL_TIMEZONE).strftime("%Y-%m-%d"),
        today_formatted=datetime.now(APP_LOCAL_TIMEZONE).strftime("%B %d, %Y"),
        success_count=stats.get("total_sent", 0),
        failed_count=stats.get("total_failed", 0),
    )


@app.get("/api/ca-certificate/stats")
def ca_certificate_stats():
    try:
        stats = repository.get_unified_history(record_type="ca_certificate", page=1, per_page=1)
        return jsonify({
            "success": True,
            "success_count": stats.get("total_sent", 0),
            "failed_count": stats.get("total_failed", 0),
        })
    except Exception as exc:
        print("CA CERTIFICATE STATS API ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.post("/api/ca-certificate/generate")
def ca_certificate_generate():
    data = request.get_json(force=True) or {}
    name = str(data.get("participant_name") or "").strip()
    email = str(data.get("participant_email") or "").strip().lower()
    date_value = str(data.get("program_date") or "").strip()
    if not name or not date_value:
        return jsonify({"success": False, "error": "Participant name and program date are required."}), 400
    try:
        formatted_date = ca_certificate_service.format_program_date(date_value)
        filename = f"{ca_certificate_service.safe_filename(name)}_CA_Certificate.pdf"
        ca_certificate_service.generate_pdf_bytes(name, date_value)
        return jsonify({
            "success": True,
            "filename": filename,
            "date": formatted_date,
            "document_id": f"CA-CERT-{uuid.uuid4().hex}",
        })
    except (ValueError, RuntimeError, FileNotFoundError) as exc:
        return jsonify({"success": False, "error": str(exc)}), 400
    except Exception as exc:
        print("CA CERTIFICATE GENERATION ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.post("/api/ca-certificate/pdf")
def ca_certificate_pdf():
    data = request.get_json(force=True) or {}
    name = str(data.get("participant_name") or "").strip()
    date_value = str(data.get("program_date") or "").strip()
    try:
        pdf_bytes, _ = ca_certificate_service.generate_pdf_bytes(name, date_value)
        filename = f"{ca_certificate_service.safe_filename(name)}_CA_Certificate.pdf"
        return send_file(
            io.BytesIO(pdf_bytes),
            mimetype="application/pdf",
            as_attachment=bool(data.get("download")),
            download_name=filename,
        )
    except (ValueError, RuntimeError, FileNotFoundError) as exc:
        return jsonify({"success": False, "error": str(exc)}), 400
    except Exception as exc:
        print("CA CERTIFICATE PDF ERROR:", repr(exc))
        return jsonify({"success": False, "error": "Certificate PDF generation failed."}), 500


@app.post("/api/ca-certificate/send-email")
def ca_certificate_send_email():
    data = request.get_json(force=True) or {}
    name = str(data.get("participant_name") or "").strip()
    email = str(data.get("participant_email") or "").strip().lower()
    date_value = str(data.get("program_date") or "").strip()
    if not name or not email or not date_value:
        return jsonify({"success": False, "error": "A generated certificate and complete participant details are required."}), 400
    filename = f"{ca_certificate_service.safe_filename(name)}_CA_Certificate.pdf"
    resend = data.get("send_again") is True or data.get("force_resend") is True
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        return jsonify({"success": False, "error": "Enter a valid email address."}), 400

    # --------------------------------------------------------
    # AUTHORITATIVE SERVER-SIDE CHECK AGAINST SUPABASE
    # --------------------------------------------------------
    existing_record = repository.get_latest_ca_certificate_by_email(email)

    if existing_record:
        existing_status = str(existing_record.get("email_status") or "").strip().lower()
        if existing_status in ("sending", "uncertain"):
            return jsonify({
                "success": False,
                "duplicate": True,
                "can_resend": False,
                "reconciliation_required": True,
                "email_status": existing_status,
                "message": "A CA Certificate delivery is already in progress or uncertain for this email address. Please reconcile the delivery status before retrying.",
                "previous_record": existing_record,
            }), 409

        if existing_status == "sent" and not resend:
            return jsonify({
                "success": False,
                "duplicate": True,
                "previously_sent": True,
                "can_resend": True,
                "email_status": "sent",
                "send_count": existing_record.get("send_count") or 1,
                "message": "A CA Certificate has already been sent to this email address. Would you like to send it again?",
                "previous_record": existing_record,
            }), 409

    # Reuse existing document_id when resending or retrying existing record so the record is updated in-place!
    if existing_record and existing_record.get("document_id"):
        document_id = str(existing_record["document_id"])
    else:
        document_id = str(data.get("document_id") or "")
        if not re.fullmatch(r"CA-CERT-[0-9a-f]{32}", document_id):
            document_id = f"CA-CERT-{uuid.uuid4().hex}"

    claim_token = str(uuid.uuid4())
    try:
        claim = repository.claim_ca_certificate_single_email(
            {
                "participant_name": name,
                "participant_email": email,
                "program_date": ca_certificate_service.format_program_date(date_value),
                "document_id": document_id,
                "pdf_filename": filename,
            },
            claim_token=claim_token,
            allow_resend=resend,
        )
    except RuntimeError as exc:
        return jsonify({"success": False, "error": str(exc)}), 503
    except Exception as exc:
        print("CA CERTIFICATE SINGLE CLAIM ERROR:", repr(exc))
        return jsonify({"success": False, "error": "Unable to claim certificate delivery safely."}), 503

    if not claim.get("claimed"):
        claim_status = str(claim.get("email_status") or "").strip().lower()
        if claim_status == "sent" and not resend:
            return jsonify({
                "success": False,
                "duplicate": True,
                "previously_sent": True,
                "can_resend": True,
                "email_status": "sent",
                "message": "A CA Certificate has already been sent to this email address. Would you like to send it again?",
            }), 409
        elif claim_status in ("sending", "uncertain"):
            return jsonify({
                "success": False,
                "duplicate": True,
                "can_resend": False,
                "reconciliation_required": True,
                "email_status": claim_status,
                "message": "A CA Certificate delivery is already in progress or uncertain for this email address. Please reconcile the delivery status before retrying.",
            }), 409
        return jsonify({
            "success": False,
            "duplicate": True,
            "can_resend": False,
            "email_status": claim.get("email_status"),
            "message": "A certificate delivery is already recorded or in progress for this email address.",
        }), 409

    existing_id = claim["record_id"]
    send_count = int(claim.get("send_count") or 1)
    send_attempted = False
    try:
        message = EmailMessage()
        message["Subject"] = "Campus Ambassador Certificate – Persevex"
        message["From"] = f"Persevex LLP <{SENDER_EMAIL}>"
        message["To"] = email
        message.set_content(
            f"Dear {name},\n\nPlease find attached your Campus Ambassador Certificate from Persevex.\n\n"
            "Congratulations on your participation.\n\nWarm regards,\nTeam Persevex"
        )
        pdf_bytes, _ = ca_certificate_service.generate_pdf_bytes(name, date_value)
        message.add_attachment(pdf_bytes, maintype="application", subtype="pdf", filename=filename)
        with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30) as smtp:
            smtp.login(SENDER_EMAIL, SENDER_PASSWORD)
            send_attempted = True
            smtp.send_message(message)
    except Exception as exc:
        email_status = "uncertain" if send_attempted else "failed"
        exc_str = str(exc)
        exc_code = getattr(exc, "smtp_code", None) or getattr(exc, "code", None)

        # Classify the SMTP failure into a user-friendly message without exposing internals.
        if exc_code == 421 or "421" in exc_str:
            user_error = (
                "The email provider temporarily blocked delivery (421 \u2013 too many failed logins). "
                "Please wait a few minutes before trying again. No email was sent."
            )
            smtp_blocked = True
            http_status = 503
        elif isinstance(exc, smtplib.SMTPAuthenticationError) or "535" in exc_str or "authentication" in exc_str.lower():
            user_error = "Email authentication failed. Please check the sender credentials and try again."
            smtp_blocked = False
            http_status = 503
        elif email_status == "uncertain":
            user_error = (
                "The email was transmitted but confirmation was not received. "
                "Please check the recipient\u2019s inbox before resending."
            )
            smtp_blocked = False
            http_status = 202
        elif isinstance(exc, (smtplib.SMTPConnectError, smtplib.SMTPServerDisconnected)):
            user_error = "Could not connect to the email server. Please check your network connection and try again."
            smtp_blocked = False
            http_status = 503
        elif isinstance(exc, smtplib.SMTPRecipientsRefused):
            user_error = "The recipient email address was rejected by the mail server."
            smtp_blocked = False
            http_status = 422
        else:
            user_error = "Email delivery failed. Please try again in a moment."
            smtp_blocked = False
            http_status = 500

        # Roll back the incremented send_count — only confirmed deliveries should count.
        failed_send_count = max(1, send_count - 1)
        try:
            persistence = repository.finish_ca_certificate_single_email(
                existing_id,
                claim_token,
                {
                    "participant_name": name, "participant_email": email,
                    "program_date": ca_certificate_service.format_program_date(date_value),
                    "document_id": document_id,
                    "pdf_filename": filename,
                    "email_status": email_status,
                    "send_count": failed_send_count,
                    "error_message": exc_str[:500],
                },
            )
        except Exception as history_exc:
            print("CA CERTIFICATE FAILED-DELIVERY HISTORY ERROR:", repr(history_exc))
            persistence = {"saved": False, "status": "not_saved"}
        return jsonify({
            "success": False,
            "error": user_error,
            "email_status": email_status,
            "smtp_blocked": smtp_blocked,
            "history_saved": persistence["saved"],
            "history_status": persistence["status"],
        }), http_status
    try:
        persistence = repository.finish_ca_certificate_single_email(
            existing_id,
            claim_token,
            {
                "participant_name": name, "participant_email": email,
                "program_date": ca_certificate_service.format_program_date(date_value),
                "document_id": document_id, "pdf_filename": filename,
                "email_status": "sent", "sent_at": datetime.now(timezone.utc).isoformat(),
                "send_count": send_count,
            },
        )
    except Exception as history_exc:
        print("CA CERTIFICATE DELIVERY HISTORY ERROR:", repr(history_exc))
        persistence = {"saved": False, "status": "not_saved", "error": str(history_exc)}
    response = {
        "success": True,
        "message": "CA Certificate sent successfully.",
        "send_count": send_count,
        "email_delivered": True,
        "history_saved": persistence["saved"],
        "history_status": persistence["status"],
    }
    if persistence["status"] == "local_only" and SUPABASE_URL and SUPABASE_KEY:
        response["warning"] = "Email delivered, but Supabase history was unavailable. A local reconciliation copy was retained; do not resend this email."
    elif persistence["status"] == "not_saved":
        response["warning"] = "Email delivered, but the history record could not be saved. Do not resend this email; reconcile this delivery before retrying."
    elif not persistence.get("saved"):
        response["warning"] = "Email delivered, but the delivery claim could not be finalized. Do not resend; reconcile the uncertain history record."
    return jsonify(response)


@app.post("/api/ca-certificate/bulk/validate")
def ca_certificate_bulk_validate():
    data = request.get_json(force=True) or {}
    names = [line.strip() for line in str(data.get("names") or "").splitlines()]
    emails = [line.strip().lower() for line in str(data.get("emails") or "").splitlines()]
    date_value = str(data.get("program_date") or "").strip()
    rows = []
    seen_emails = set()
    total_rows = max(len(names), len(emails))
    historical_sent = repository.get_previous_sent_ca_certificates(emails)
    for index in range(total_rows):
        name = names[index] if index < len(names) else ""
        email = emails[index] if index < len(emails) else ""
        errors = []
        if not name: errors.append("Missing participant name")
        if not email:
            errors.append("Missing email address")
        elif not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            errors.append("Invalid email address")
        duplicate = bool(email and not errors and email in seen_emails)
        if duplicate:
            errors.append("Duplicate email address")
        elif email and not errors and email in historical_sent:
            errors.append("Email already received a CA Certificate")
            duplicate = True
        if email and not errors:
            seen_emails.add(email)
        rows.append({
            "row": index + 1,
            "participant_name": name,
            "participant_email": email,
            "valid": not errors,
            "duplicate": duplicate,
            "errors": errors,
            "status": "duplicate" if duplicate else ("valid" if not errors else "invalid"),
            "remarks": "; ".join(errors) if errors else "Ready for generation"
        })
    try:
        ca_certificate_service.format_program_date(date_value)
    except ValueError as exc:
        for row in rows:
            if row["valid"]:
                row["errors"].append(str(exc))
                row["remarks"] = str(exc)
                row["valid"] = False
                row["status"] = "invalid"
    duplicate_count = sum(1 for row in rows if row["duplicate"])
    invalid_count = sum(1 for row in rows if not row["valid"] and not row["duplicate"])
    return jsonify({
        "success": True,
        "rows": rows,
        "total_rows": total_rows,
        "valid_count": sum(1 for row in rows if row["valid"]),
        "invalid_count": invalid_count,
        "duplicate_count": duplicate_count
    })


@app.post("/api/ca-certificate/bulk/jobs")
def ca_certificate_bulk_create_job():
    data = request.get_json(force=True) or {}
    date_value = str(data.get("program_date") or "").strip()
    rows = data.get("rows") or []
    try:
        formatted_date = ca_certificate_service.format_program_date(date_value)
        items = []
        seen = set()
        for row in rows:
            if not row.get("valid"):
                continue
            email = str(row.get("participant_email") or "").strip().lower()
            if email in seen:
                continue
            seen.add(email)
            items.append({
                "source_row": row.get("row"),
                "participant_name": str(row.get("participant_name") or "").strip(),
                "participant_email": email,
            })
        if not items:
            return jsonify({"success": False, "error": "No valid rows are available for processing."}), 400
        job = repository.create_ca_certificate_job(
            formatted_date,
            items,
            skipped_count=len(rows) - len(items),
        )
        return jsonify({"success": True, "job": job, "count": len(items)})
    except (ValueError, RuntimeError) as exc:
        return jsonify({"success": False, "error": str(exc)}), 400


@app.get("/api/ca-certificate/bulk/jobs/<job_id>")
def ca_certificate_bulk_job_status(job_id):
    try:
        job = repository.get_ca_certificate_job(job_id)
        if not job:
            return jsonify({"success": False, "error": "CA Certificate job not found."}), 404
        job["items"] = repository.get_ca_certificate_job_items(job_id)
        return jsonify({"success": True, "job": job})
    except RuntimeError as exc:
        return jsonify({"success": False, "error": str(exc)}), 503


@app.post("/api/ca-certificate/bulk/jobs/<job_id>/send")
def ca_certificate_bulk_send_job(job_id):
    data = request.get_json(force=True) or {}
    try:
        items = repository.claim_ca_certificate_email_items(
            job_id, 1, retry_failed=data.get("retry") is True
        )
        sent = failed = 0
        uncertain = 0
        failed_items = []
        history_warnings = []
        for item in items[:1]:
            document_id = item.get("document_id") or f"CA-CERT-{uuid.uuid4().hex}"
            filename = f"{ca_certificate_service.safe_filename(item['participant_name'])}_CA_Certificate.pdf"
            send_attempted = False
            try:
                claim_record = repository.save_ca_certificate_record({
                    "participant_name": item["participant_name"],
                    "participant_email": item["participant_email"],
                    "program_date": item["program_date"],
                    "document_id": document_id,
                    "pdf_filename": filename,
                    "email_status": "sending",
                    "send_count": int(item.get("send_count") or 0),
                })
                if not claim_record.get("supabase_saved"):
                    raise RuntimeError("Could not persist the CA Certificate delivery claim.")
            except Exception as exc:
                history_warnings.append(document_id)
                repository.update_ca_certificate_job_item(item["id"], {
                    "document_id": document_id,
                    "email_status": "failed",
                    "attempt_count": int(item.get("attempt_count") or 0) + 1,
                    "error_message": str(exc),
                    "updated_at": datetime.now(timezone.utc).isoformat(),
                }, claim_token=item.get("claim_token"))
                failed += 1
                failed_items.append({
                    "participant_name": item["participant_name"],
                    "participant_email": item["participant_email"],
                    "program_date": item["program_date"],
                    "error": str(exc),
                })
                continue
            try:
                pdf_bytes, formatted_date = ca_certificate_service.generate_pdf_bytes(
                    item["participant_name"], item["program_date"]
                )
                message = EmailMessage()
                message["Subject"] = "Campus Ambassador Certificate – Persevex"
                message["From"] = f"Persevex LLP <{SENDER_EMAIL}>"
                message["To"] = item["participant_email"]
                message.set_content(
                    f"Dear {item['participant_name']},\n\nPlease find attached your Campus Ambassador Certificate from Persevex.\n\n"
                    "Congratulations on your participation.\n\nWarm regards,\nTeam Persevex"
                )
                message.add_attachment(pdf_bytes, maintype="application", subtype="pdf", filename=filename)
                with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30) as smtp:
                    smtp.login(SENDER_EMAIL, SENDER_PASSWORD)
                    send_attempted = True
                    smtp.send_message(message)
            except Exception as exc:
                email_status = "uncertain" if send_attempted else "failed"
                try:
                    persistence = repository.save_ca_certificate_record({
                        "participant_name": item["participant_name"],
                        "participant_email": item["participant_email"],
                        "program_date": item["program_date"],
                        "document_id": document_id,
                        "pdf_filename": filename,
                        "email_status": email_status,
                        "send_count": int(item.get("send_count") or 0),
                        "error_message": str(exc),
                    })
                    if not persistence.get("supabase_saved"):
                        history_warnings.append(document_id)
                except Exception as history_exc:
                    print("CA CERTIFICATE DELIVERY HISTORY ERROR:", repr(history_exc))
                    history_warnings.append(document_id)
                repository.update_ca_certificate_job_item(item["id"], {
                    "document_id": document_id,
                    "email_status": email_status,
                    "attempt_count": int(item.get("attempt_count") or 0) + 1,
                    "error_message": str(exc),
                    "updated_at": datetime.now(timezone.utc).isoformat(),
                }, claim_token=item.get("claim_token"))
                if email_status == "uncertain":
                    uncertain += 1
                else:
                    failed += 1
                    failed_items.append({
                        "participant_name": item["participant_name"],
                        "participant_email": item["participant_email"],
                        "program_date": item["program_date"],
                        "error": str(exc),
                    })
                continue

            try:
                persistence = repository.save_ca_certificate_record({
                    "participant_name": item["participant_name"],
                    "participant_email": item["participant_email"],
                    "program_date": formatted_date,
                    "document_id": document_id,
                    "pdf_filename": filename,
                    "email_status": "sent",
                    "sent_at": datetime.now(timezone.utc).isoformat(),
                    "send_count": int(item.get("send_count") or 0) + 1,
                })
                if not persistence.get("supabase_saved"):
                    history_warnings.append(document_id)
            except Exception as history_exc:
                print("CA CERTIFICATE DELIVERY HISTORY ERROR:", repr(history_exc))
                history_warnings.append(document_id)
            repository.update_ca_certificate_job_item(item["id"], {
                "document_id": document_id,
                "email_status": "sent",
                "send_count": int(item.get("send_count") or 0) + 1,
                "sent_at": datetime.now(timezone.utc).isoformat(),
                "error_message": None,
                "updated_at": datetime.now(timezone.utc).isoformat(),
            }, claim_token=item.get("claim_token"))
            sent += 1
        counts = repository.reconcile_ca_certificate_job(job_id)
        eligible_statuses = ("failed",) if data.get("retry") is True else ("pending",)
        remaining_items = repository.get_ca_certificate_job_items(job_id, eligible_statuses)
        return jsonify({
            "success": True,
            "claimed": len(items),
            "sent": sent,
            "failed": failed,
            "uncertain": uncertain,
            "failed_items": failed_items,
            "remaining": len(remaining_items),
            "counts": counts,
            "history_warnings": history_warnings,
        })
    except RuntimeError as exc:
        return jsonify({"success": False, "error": str(exc)}), 503


@app.post("/api/ca-certificate/bulk/jobs/<job_id>/process")
def ca_certificate_bulk_process_job(job_id):
    data = request.get_json(force=True) or {}
    try:
        job = repository.get_ca_certificate_job(job_id)
        if not job:
            return jsonify({"success": False, "error": "CA Certificate job not found."}), 404
        batch = repository.claim_ca_certificate_job_items(job_id, data.get("batch_size", 10))
        if not batch:
            counts = repository.reconcile_ca_certificate_job(job_id)
            return jsonify({"success": True, "job_id": job_id, "processed": 0, "remaining": 0, "complete": counts["status"] != "running", "counts": counts})
        processed = generated = failed = 0
        for item in batch:
            processed += 1
            name = item["participant_name"]
            email = item["participant_email"]
            filename = f"{ca_certificate_service.safe_filename(name)}_CA_Certificate.pdf"
            try:
                pdf_bytes, formatted_date = ca_certificate_service.generate_pdf_bytes(name, item["program_date"])
                document_id = f"CA-CERT-{uuid.uuid4().hex}"
                repository.update_ca_certificate_job_item(item["id"], {
                    "generation_status": "generated",
                    "email_status": "pending",
                    "document_id": document_id,
                    "updated_at": datetime.now(timezone.utc).isoformat(),
                }, claim_token=item.get("claim_token"))
                generated += 1
            except Exception as exc:
                failed += 1
                repository.update_ca_certificate_job_item(item["id"], {
                    "generation_status": "failed",
                    "email_status": "failed",
                    "attempt_count": int(item.get("attempt_count") or 0) + 1,
                    "error_message": str(exc),
                    "updated_at": datetime.now(timezone.utc).isoformat(),
                }, claim_token=item.get("claim_token"))
        counts = repository.reconcile_ca_certificate_job(job_id)
        return jsonify({"success": True, "job_id": job_id, "processed": processed, "generated": generated, "failed": failed, "counts": counts, "complete": len(batch) < min(max(int(data.get("batch_size", 10)), 1), 25)})
    except RuntimeError as exc:
        return jsonify({"success": False, "error": str(exc)}), 503


@app.post("/api/ca-certificate/bulk/parse-upload")
def ca_certificate_bulk_parse_upload():
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        return jsonify({"success": False, "error": "CSV or Excel file is required."}), 400
    try:
        file_bytes = uploaded.read()
        rows = parse_ca_file(file_bytes, uploaded.filename, preserve_blank_rows=True)
        if not rows.get("success"):
            return jsonify(rows), 400
        rows["file_size_bytes"] = len(file_bytes)
        rows["file_size_label"] = (
            f"{len(file_bytes) / 1024:.1f} KB"
            if len(file_bytes) < 1024 * 1024
            else f"{len(file_bytes) / (1024 * 1024):.1f} MB"
        )
        return jsonify(rows)
    except Exception as exc:
        return jsonify({"success": False, "error": str(exc)}), 400


@app.post("/api/ca-certificate/bulk/generate")
def ca_certificate_bulk_generate():
    data = request.get_json(force=True) or {}
    rows = data.get("rows") or []
    date_value = str(data.get("program_date") or "").strip()
    generated = []
    errors = []
    for index, row in enumerate(rows, 1):
        name = str(row.get("participant_name") or row.get("student_name") or "").strip()
        email = str(row.get("participant_email") or row.get("student_email") or "").strip().lower()
        if not name or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            errors.append({"row": index, "error": "Invalid participant name or email."})
            continue
        try:
            filename = f"{ca_certificate_service.safe_filename(name)}_CA_Certificate.pdf"
            ca_certificate_service.generate_pdf_bytes(name, date_value)
            generated.append({"row": index, "participant_name": name, "participant_email": email, "filename": filename})
        except Exception as exc:
            errors.append({"row": index, "error": str(exc)})
    return jsonify({"success": True, "generated": generated, "errors": errors})


@app.post("/api/ca-certificate/bulk/download")
def ca_certificate_bulk_download():
    data = request.get_json(force=True) or {}
    certificates = data.get("certificates") or []
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
        for certificate in certificates[:25]:
            name = str(certificate.get("participant_name") or "").strip()
            date_value = str(certificate.get("program_date") or "").strip()
            if not name or not date_value:
                continue
            pdf_bytes, _ = ca_certificate_service.generate_pdf_bytes(name, date_value)
            filename = f"{ca_certificate_service.safe_filename(name)}_CA_Certificate.pdf"
            bundle.writestr(filename, pdf_bytes)
    archive.seek(0)
    return Response(
        archive.read(),
        mimetype="application/zip",
        headers={"Content-Disposition": "attachment; filename=Persevex_CA_Certificates.zip"},
    )

@app.get("/certificate")
def certificate_page():
    try:
        history_data = repository.get_unified_history(record_type="certificate", page=1, per_page=1)
        success_count = history_data.get("total_sent", 0)
        failed_count = history_data.get("total_failed", 0)
    except Exception as exc:
        print("CERTIFICATE STATS COUNT ERROR:", repr(exc))
        success_count = 0
        failed_count = 0

    return render_template(
        "certificate.html",
        success_count=success_count,
        failed_count=failed_count,
    )


@app.get("/api/certificate/status-counts")
def certificate_status_counts():
    try:
        history_data = repository.get_unified_history(record_type="certificate", page=1, per_page=1)
        return jsonify({
            "success": True,
            "total_count": history_data.get("total_certificates", 0),
            "success_count": history_data.get("total_sent", 0),
            "failed_count": history_data.get("total_failed", 0),
        })
    except Exception as exc:
        print("CERTIFICATE STATUS COUNTS ERROR:", repr(exc))
        return jsonify({
            "success": False,
            "total_count": 0,
            "success_count": 0,
            "failed_count": 0,
        }), 500


@app.post("/api/certificate/generate")
def certificate_generate():
    try:
        data = request.get_json(force=True) or {}
        if not data:
            return jsonify({"error": "No certificate data received."}), 400

        student_name = str(data.get("student_name", "")).strip()
        student_email = str(data.get("student_email", "")).strip().lower()
        domain = str(data.get("domain") or data.get("internship_domain") or "").strip()
        start_date = str(data.get("start_date", "")).strip()
        end_date = str(data.get("end_date", "")).strip()
        issued_date = str(data.get("issued_date", "")).strip()

        if not student_name or not student_email or not domain or not start_date or not end_date or not issued_date:
            return jsonify({
                "error": "Please fill all required fields: Student Name, Email, Domain, Start Date, End Date, and Issued Date."
            }), 400

        try:
            start_date_obj = datetime.strptime(start_date, "%Y-%m-%d")
            end_date_obj = datetime.strptime(end_date, "%Y-%m-%d")
            if "-" in issued_date and len(issued_date.split("-")[0]) == 4:
                datetime.strptime(issued_date, "%Y-%m-%d")
        except ValueError:
            return jsonify({"error": "Invalid date format. Expected YYYY-MM-DD."}), 400

        if end_date_obj < start_date_obj:
            return jsonify({"error": "End date cannot be before start date."}), 400

        data["domain"] = domain
        data["internship_domain"] = domain

        formatted_start = certificate_service.format_display_date(start_date)
        formatted_end = certificate_service.format_display_date(end_date)
        formatted_issued = certificate_service.format_display_date(issued_date)

        # Idempotent lookup: Check for existing active certificate with exact candidate identity
        existing_cert = repository.find_active_certificate(
            email=student_email,
            domain=domain,
            start_date=formatted_start,
            end_date=formatted_end,
        )

        if existing_cert and str(existing_cert.get("certificate_status") or "").lower() != "revoked":
            cert_id = str(existing_cert.get("certificate_id") or "").strip()
            student_name_safe = certificate_service.safe_filename(existing_cert.get("student_name") or student_name)
            filename = str(existing_cert.get("pdf_filename") or f"{student_name_safe}_Certificate.pdf").strip()
            preview_params = urllib.parse.urlencode({
                "name": existing_cert.get("student_name") or student_name,
                "email": student_email,
                "domain": existing_cert.get("internship_domain") or domain,
                "start": existing_cert.get("start_date") or formatted_start,
                "end": existing_cert.get("end_date") or formatted_end,
                "issued": existing_cert.get("issued_date") or formatted_issued,
            })
            preview_url = f"/verify/{cert_id}/preview?{preview_params}"
            download_url = f"/verify/{cert_id}/download?{preview_params}"

            return jsonify({
                "success": True,
                "filename": filename,
                "url": preview_url,
                "download_url": download_url,
                "certificate_id": cert_id,
                "data": {
                    "student_name": existing_cert.get("student_name") or student_name,
                    "student_email": student_email,
                    "domain": existing_cert.get("internship_domain") or domain,
                    "start_date": existing_cert.get("start_date") or formatted_start,
                    "end_date": existing_cert.get("end_date") or formatted_end,
                    "issued_date": existing_cert.get("issued_date") or formatted_issued,
                    "certificate_id": cert_id,
                }
            })

        # 2. Check if student already has a previous sent certificate -> reuse that certificate_id
        previous_sent_cert = repository.get_previous_sent_certificate_by_email(student_email)
        if previous_sent_cert and str(previous_sent_cert.get("certificate_status") or "").lower() != "revoked":
            cert_id = str(previous_sent_cert.get("certificate_id") or "").strip()
        else:
            cert_id = certificate_service.generate_certificate_id(supabase)

        data["certificate_id"] = cert_id
        base_url = certificate_service.get_public_base_url(request)

        filename = certificate_service.generate_certificate_pdf(
            data=data,
            output_dir=GENERATED_DIR,
            base_url=base_url,
        )

        preview_params = urllib.parse.urlencode({
            "name": student_name,
            "email": student_email,
            "domain": domain,
            "start": formatted_start,
            "end": formatted_end,
            "issued": formatted_issued,
        })
        preview_url = f"/verify/{cert_id}/preview?{preview_params}"
        download_url = f"/verify/{cert_id}/download?{preview_params}"

        return jsonify({
            "success": True,
            "filename": filename,
            "url": preview_url,
            "download_url": download_url,
            "certificate_id": cert_id,
            "data": {
                "student_name": student_name,
                "student_email": student_email,
                "domain": domain,
                "start_date": formatted_start,
                "end_date": formatted_end,
                "issued_date": formatted_issued,
                "certificate_id": cert_id,
            }
        })

    except Exception as exc:
        print("CERTIFICATE GENERATE API ERROR:", repr(exc))
        return jsonify({"error": str(exc)}), 500


@app.post("/api/certificate/send-email")
def certificate_send_email():
    try:
        data = request.get_json(force=True) or {}
        filename = str(data.get("filename") or "").strip()
        recipient = str(data.get("student_email") or "").strip().lower()
        student_name = str(data.get("student_name") or "Student").strip()
        domain = str(data.get("domain") or data.get("internship_domain") or "Internship").strip()
        cert_id = str(data.get("certificate_id") or "").strip()
        start_date = str(data.get("start_date") or "").strip()
        end_date = str(data.get("end_date") or "").strip()
        issued_date = str(data.get("issued_date") or "").strip()
        send_again = data.get("send_again") is True

        if not recipient:
            return jsonify({"success": False, "error": "Student email address is required."}), 400

        # Resend detection based on normalized recipient email ONLY
        previous_sent_record = None
        if not send_again:
            previous_sent_record = repository.get_previous_sent_certificate_by_email(recipient)
            if previous_sent_record and str(previous_sent_record.get("email_status") or "").strip().lower() == "sent":
                return jsonify({
                    "success": False,
                    "duplicate": True,
                    "message": "This email has already received a certificate. Do you want to send it again?",
                    "previous_record": previous_sent_record,
                }), 409
        else:
            previous_sent_record = repository.get_previous_sent_certificate_by_email(recipient)

        is_resend = bool(send_again and previous_sent_record and str(previous_sent_record.get("email_status") or "").strip().lower() == "sent")
        target_record_id = (data.get("id") or data.get("existing_id") or (previous_sent_record.get("id") if previous_sent_record else None)) if is_resend else None

        # Persistent Certificate ID: if resend, MUST preserve the existing certificate_id
        if is_resend and previous_sent_record:
            cert_id = str(previous_sent_record.get("certificate_id") or cert_id).strip()
            prev_count = int(previous_sent_record.get("send_count") or 1)
            current_send_count = prev_count + 1
        else:
            current_send_count = 1
            if not cert_id:
                cert_id = certificate_service.generate_certificate_id(supabase)

        current_record = None
        if cert_id:
            current_record = repository.get_certificate_by_id(cert_id)

        cert_data = {
            "student_name": student_name,
            "student_email": recipient,
            "domain": domain,
            "start_date": start_date or (current_record.get("start_date") if current_record else (previous_sent_record.get("start_date") if previous_sent_record else "")),
            "end_date": end_date or (current_record.get("end_date") if current_record else (previous_sent_record.get("end_date") if previous_sent_record else "")),
            "issued_date": issued_date or (current_record.get("issued_date") if current_record else (previous_sent_record.get("issued_date") if previous_sent_record else "")),
            "certificate_id": cert_id,
            "template_version": data.get("template_version") or (current_record.get("template_version") if current_record else "v1"),
        }

        # Regenerate/ensure PDF has the exact cert_id, QR code, and current details
        base_url = certificate_service.get_public_base_url(request)
        safe_file = certificate_service.generate_certificate_pdf(
            data=cert_data,
            output_dir=GENERATED_DIR,
            base_url=base_url,
            template_version=cert_data.get("template_version", "v1")
        )
        pdf_path = GENERATED_DIR / safe_file

        sent_time = datetime.now(timezone.utc).isoformat()
        try:
            certificate_service.send_certificate_email(
                data=cert_data,
                pdf_path=pdf_path,
                filename=safe_file,
            )
        except Exception as smtp_exc:
            error_msg = str(smtp_exc)
            print("CERTIFICATE SMTP SEND ERROR:", repr(smtp_exc))

            fail_record = {
                "certificate_id": cert_data["certificate_id"],
                "student_name": student_name,
                "student_email": recipient,
                "internship_domain": domain,
                "start_date": cert_data["start_date"],
                "end_date": cert_data["end_date"],
                "issued_date": cert_data["issued_date"],
                "email_status": "failed",
                "sent_at": None,
                "send_count": current_send_count,
                "pdf_filename": safe_file,
                "error_message": error_msg,
                "template_version": cert_data.get("template_version", "v1"),
            }
            if is_resend and target_record_id:
                fail_record["id"] = target_record_id
                fail_record["send_again"] = True
            repository.save_certificate_record(fail_record, existing_id=target_record_id if is_resend else None)

            return jsonify({
                "success": False,
                "error": f"Failed to send email: {error_msg}",
                "email_status": "failed",
            }), 500

        # Success: update in-place if resend, or insert if first send
        success_record = {
            "certificate_id": cert_data["certificate_id"],
            "student_name": student_name,
            "student_email": recipient,
            "internship_domain": domain,
            "start_date": cert_data["start_date"],
            "end_date": cert_data["end_date"],
            "issued_date": cert_data["issued_date"],
            "email_status": "sent",
            "sent_at": sent_time,
            "send_count": current_send_count,
            "pdf_filename": safe_file,
            "error_message": None,
            "template_version": cert_data.get("template_version", "v1"),
        }
        if is_resend and target_record_id:
            success_record["id"] = target_record_id
            success_record["send_again"] = True

        repository.save_certificate_record(success_record, existing_id=target_record_id if is_resend else None)

        return jsonify({
            "success": True,
            "message": "Certificate sent successfully.",
            "email_status": "sent",
            "send_count": current_send_count,
            "certificate_id": cert_data["certificate_id"],
        })

    except Exception as exc:
        print("CERTIFICATE SEND EMAIL ROUTE ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


# ============================================================
# BULK CERTIFICATE API ENDPOINTS
# ============================================================

@app.post("/api/certificate/bulk-check-duplicates")
def certificate_bulk_check_duplicates():
    """
    Check uploaded certificate batch for historical sent recipients.
    Considers ONLY records where email_status = 'sent'.
    Normalizes candidate emails and database emails using email.strip().lower().
    Authoritative source: Supabase (falls back to local SQLite on exception/unconfigured).
    Returns list of duplicate email strings and mapping of details.
    """
    try:
        data = request.get_json(force=True) or {}
        records = data.get("records") or data.get("emails") or []
        if not records:
            return jsonify({
                "success": True,
                "duplicate_count": 0,
                "duplicates": [],
                "duplicate_map": {}
            })

        duplicate_map = repository.find_existing_certificate_emails_batch(records)
        duplicate_emails = list(duplicate_map.keys())

        return jsonify({
            "success": True,
            "duplicate_count": len(duplicate_emails),
            "duplicates": duplicate_emails,
            "duplicate_map": duplicate_map
        })
    except Exception as exc:
        print("BULK CERTIFICATE CHECK DUPLICATES ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.post("/api/certificate/bulk/validate")
def bulk_certificate_validate():
    try:
        file = request.files.get("file")
        if not file or not file.filename:
            return jsonify({"success": False, "error": "No file uploaded."}), 400

        filename = file.filename.lower()
        file_bytes = file.read()

        if filename.endswith(".xlsx") or filename.endswith(".xls"):
            parsed_rows = bulk_certificate_service.parse_excel_content(file_bytes)
        elif filename.endswith(".csv") or filename.endswith(".txt"):
            parsed_rows = bulk_certificate_service.parse_csv_content(file_bytes)
        else:
            return jsonify({"success": False, "error": "Unsupported file type. Please upload CSV or Excel (.xlsx)."}), 400

        if not parsed_rows:
            return jsonify({"success": False, "error": "The uploaded spreadsheet is empty or has invalid headers."}), 400

        validation_result = bulk_certificate_service.validate_bulk_records(parsed_rows)
        return jsonify({
            "success": True,
            **validation_result
        })

    except Exception as exc:
        print("BULK VALIDATE ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.post("/api/certificate/preview-sample")
def bulk_certificate_preview_sample():
    """
    Renders an in-memory preview of an individual candidate's certificate from uploaded spreadsheet data.
    Does NOT write to database, does NOT send emails, does NOT create permanent files.
    """
    try:
        data = request.get_json(force=True) or {}
        student_name = str(data.get("student_name") or "").strip()
        domain = str(data.get("domain") or data.get("internship_domain") or "").strip()
        if not student_name:
            return jsonify({"success": False, "error": "Student Name is required for preview."}), 400

        base_url = certificate_service.get_public_base_url(request)
        preview_data = {
            "student_name": student_name,
            "domain": domain or "Internship Domain",
            "start_date": data.get("start_date") or "",
            "end_date": data.get("end_date") or "",
            "issued_date": data.get("issued_date") or "",
            "certificate_id": "PXL-CERT-PREVIEW",
        }
        preview_img = certificate_service.render_certificate_preview_image(preview_data, base_url=base_url)
        return jsonify({
            "success": True,
            "preview_image": preview_img,
            "student_name": student_name,
            "domain": domain
        })
    except Exception as exc:
        print("BULK CERTIFICATE PREVIEW SAMPLE ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.post("/api/certificate/bulk/generate")
def bulk_certificate_generate():
    try:
        data = request.get_json(force=True) or {}
        rows = data.get("rows") or []
        if not rows:
            return jsonify({"success": False, "error": "No valid rows provided for certificate generation."}), 400

        base_url = certificate_service.get_public_base_url(request)
        result = bulk_certificate_service.generate_bulk_certificates(
            valid_rows=rows,
            output_dir=GENERATED_DIR,
            base_url=base_url,
            supabase_client=supabase
        )

        return jsonify({
            "success": True,
            **result
        })

    except Exception as exc:
        print("BULK GENERATE ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.post("/api/certificate/bulk/send-email")
def bulk_certificate_send_email():
    try:
        data = request.get_json(force=True) or {}
        items = data.get("items") or []
        skip_duplicate_emails = data.get("skip_duplicate_emails") or []
        send_again_all = bool(data.get("send_again_all", False))

        if not items:
            return jsonify({"success": False, "error": "No certificates selected for email dispatch."}), 400

        result = bulk_certificate_service.send_bulk_certificate_emails(
            certificate_items=items,
            output_dir=GENERATED_DIR,
            batch_size=5,
            delay_seconds=0.4,
            skip_duplicate_emails=skip_duplicate_emails,
            send_again_all=send_again_all
        )

        return jsonify({
            "success": True,
            **result
        })

    except Exception as exc:
        print("BULK SEND EMAIL ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.get("/api/certificate/bulk/sample-template")
def bulk_certificate_sample_template():
    try:
        fmt = (request.args.get("format") or "csv").strip().lower()
        if fmt in ("excel", "xlsx"):
            excel_bytes = bulk_certificate_service.get_sample_excel_template()
            if excel_bytes:
                return Response(
                    excel_bytes,
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={
                        "Content-Disposition": 'attachment; filename="Persevex_Certificate_Bulk_Sample.xlsx"',
                        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    }
                )

        csv_content = bulk_certificate_service.get_sample_csv_template()
        return Response(
            csv_content,
            mimetype="text/csv",
            headers={
                "Content-Disposition": 'attachment; filename="Persevex_Certificate_Bulk_Sample.csv"',
                "Content-Type": "text/csv; charset=utf-8",
            }
        )
    except Exception as exc:
        print("SAMPLE TEMPLATE ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.get("/certificate-history")
def certificate_history_legacy():
    """Redirect legacy certificate-history route to Unified History."""
    return redirect("/history?record_type=certificate", code=302)


@app.get("/certificate-history/export")
def certificate_history_export_legacy():
    return redirect("/history/export?record_type=certificate", code=302)


@app.delete("/api/certificate-history/<path:record_id>")
def delete_certificate_record_legacy(record_id):
    try:
        repository.delete_certificate_record(record_id)
        return jsonify({
            "success": True,
            "message": "Certificate record deleted successfully."
        })
    except Exception as exc:
        print("DELETE CERTIFICATE RECORD ERROR:", repr(exc))
        return jsonify({"error": str(exc)}), 500



# ============================================================
# PUBLIC CERTIFICATE VERIFICATION ROUTE (NO LOGIN REQUIRED)
# ============================================================

PERSISTENT_CERT_DIR = Path(tempfile.gettempdir()) / "persevex_certificates"
try:
    PERSISTENT_CERT_DIR.mkdir(parents=True, exist_ok=True)
except OSError:
    pass


@app.get("/verify/<certificate_id>")
def verify_certificate(certificate_id):
    """
    Publicly accessible endpoint for validating authentic Persevex certificates.
    Scanned via QR Code on mobile or opened in browser.
    Differentiates 4 states:
      1. Verified Active (200)
      2. Revoked (200 / 410)
      3. Service Unavailable (503)
      4. Not Found (404)
    """
    cert_id = str(certificate_id).strip()
    record, db_error = certificate_service.db_get_certificate_by_id(cert_id, supabase, return_error_detail=True)

    base_url = certificate_service.get_public_base_url(request)
    verification_url = certificate_service.get_verification_url(cert_id, request)

    if record:
        status = str(record.get("certificate_status") or "active").strip().lower()
        if status == "revoked":
            return render_template(
                "verify.html",
                valid=False,
                status="revoked",
                certificate=record,
                certificate_id=cert_id,
                base_url=base_url,
                verification_url=verification_url,
                revocation_reason=record.get("error_message") or "This certificate has been revoked by the issuing authority.",
            ), 200

        thumbnail_url = certificate_service.get_certificate_thumbnail_url(cert_id, request)
        pdf_preview_url = certificate_service.get_certificate_preview_url(cert_id, request)
        download_url = certificate_service.get_certificate_download_url(cert_id, request)

        linkedin_text, linkedin_share_url = certificate_service.get_linkedin_share_data(record, request)
        whatsapp_text, whatsapp_share_url = certificate_service.get_whatsapp_share_data(record, request)

        return render_template(
            "verify.html",
            valid=True,
            status="verified",
            certificate=record,
            certificate_id=cert_id,
            base_url=base_url,
            verification_url=verification_url,
            thumbnail_url=thumbnail_url,
            pdf_preview_url=pdf_preview_url,
            download_url=download_url,
            linkedin_text=linkedin_text,
            whatsapp_text=whatsapp_text,
            linkedin_share_url=linkedin_share_url,
            whatsapp_share_url=whatsapp_share_url,
        ), 200

    # If Supabase experienced a connection error and SQLite didn't have it
    if db_error and db_error != "empty_id":
        return render_template(
            "verify.html",
            valid=False,
            status="service_unavailable",
            certificate_id=cert_id,
            base_url=base_url,
            verification_url=verification_url,
        ), 503

    return render_template(
        "verify.html",
        valid=False,
        status="not_found",
        certificate_id=cert_id,
        base_url=base_url,
        verification_url=verification_url,
    ), 404


@app.post("/api/certificate/revoke/<path:certificate_id>")
def revoke_certificate_api(certificate_id):
    """
    Administrator endpoint to explicitly revoke a certificate.
    """
    try:
        data = request.get_json(silent=True) or {}
        reason = data.get("reason") or "Revoked by administrator"
        repository.revoke_certificate_record(certificate_id, reason=reason)
        return jsonify({
            "success": True,
            "message": f"Certificate {certificate_id} has been revoked successfully."
        })
    except Exception as exc:
        print("REVOKE CERTIFICATE API ERROR:", repr(exc))
        return jsonify({"success": False, "error": str(exc)}), 500


@app.get("/verify/<certificate_id>/download")
def verify_download_certificate(certificate_id):
    """
    Publicly accessible endpoint to download the authentic certificate PDF.
    Generates and streams directly from memory on demand from Supabase certificate record or query params.
    NO PERMANENT LOCAL PDF STORAGE REQUIRED.
    """
    cert_id = str(certificate_id).strip()
    record = certificate_service.db_get_certificate_by_id(cert_id, supabase)

    if record:
        if str(record.get("certificate_status") or "").lower() == "revoked":
            return jsonify({"error": "This certificate has been revoked and is unavailable for download."}), 410
        data = {
            "student_name": record.get("student_name", ""),
            "student_email": record.get("student_email", ""),
            "domain": record.get("internship_domain", ""),
            "start_date": record.get("start_date", ""),
            "end_date": record.get("end_date", ""),
            "issued_date": record.get("issued_date", ""),
            "certificate_id": cert_id,
        }
        template_version = record.get("template_version") or "v1"
    else:
        req_name = request.args.get("name") or request.args.get("student_name")
        req_domain = request.args.get("domain") or request.args.get("internship_domain")
        if req_name or req_domain:
            data = {
                "student_name": req_name or "Student",
                "student_email": request.args.get("email") or request.args.get("student_email") or "",
                "domain": req_domain or "Internship",
                "start_date": request.args.get("start") or request.args.get("start_date") or "",
                "end_date": request.args.get("end") or request.args.get("end_date") or "",
                "issued_date": request.args.get("issued") or request.args.get("issued_date") or "",
                "certificate_id": cert_id,
            }
            template_version = request.args.get("version") or "v1"
        else:
            matched_files = list(GENERATED_DIR.glob(f"*{cert_id}*.pdf"))
            if matched_files:
                return send_from_directory(GENERATED_DIR, matched_files[0].name, as_attachment=True)
            return jsonify({"error": "Certificate record not found."}), 404

    try:
        base_url = certificate_service.get_public_base_url(request)

        pdf_bytes, filename = certificate_service.generate_certificate_pdf_bytes(
            data=data,
            base_url=base_url,
            template_version=template_version,
        )

        student_name_safe = certificate_service.safe_filename(data.get("student_name", "Student"))
        download_name = f"Persevex_Certificate_{student_name_safe}_{cert_id}.pdf"

        return Response(
            pdf_bytes,
            mimetype="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{download_name}"',
                "Content-Type": "application/pdf",
                "Cache-Control": "public, max-age=3600",
            },
        )
    except Exception as exc:
        print("CERTIFICATE DOWNLOAD ERROR:", repr(exc))
        return jsonify({"error": "Certificate document could not be generated at this time."}), 500


@app.get("/verify/<certificate_id>/pdf")
@app.get("/verify/<certificate_id>/preview")
def verify_preview_certificate(certificate_id):
    """
    Publicly accessible endpoint to stream the certificate PDF for embedded preview.
    Generates and streams directly from memory on demand from Supabase certificate record or query params.
    NO PERMANENT LOCAL PDF STORAGE REQUIRED.
    """
    cert_id = str(certificate_id).strip()
    record = certificate_service.db_get_certificate_by_id(cert_id, supabase)

    if record:
        if str(record.get("certificate_status") or "").lower() == "revoked":
            return jsonify({"error": "This certificate has been revoked and is unavailable for preview."}), 410
        data = {
            "student_name": record.get("student_name", ""),
            "student_email": record.get("student_email", ""),
            "domain": record.get("internship_domain", ""),
            "start_date": record.get("start_date", ""),
            "end_date": record.get("end_date", ""),
            "issued_date": record.get("issued_date", ""),
            "certificate_id": cert_id,
        }
        template_version = record.get("template_version") or "v1"
    else:
        req_name = request.args.get("name") or request.args.get("student_name")
        req_domain = request.args.get("domain") or request.args.get("internship_domain")
        if req_name or req_domain:
            data = {
                "student_name": req_name or "Student",
                "student_email": request.args.get("email") or request.args.get("student_email") or "",
                "domain": req_domain or "Internship",
                "start_date": request.args.get("start") or request.args.get("start_date") or "",
                "end_date": request.args.get("end") or request.args.get("end_date") or "",
                "issued_date": request.args.get("issued") or request.args.get("issued_date") or "",
                "certificate_id": cert_id,
            }
            template_version = request.args.get("version") or "v1"
        else:
            matched_files = list(GENERATED_DIR.glob(f"*{cert_id}*.pdf"))
            if matched_files:
                return send_from_directory(GENERATED_DIR, matched_files[0].name, as_attachment=False)
            return jsonify({"error": "Certificate record not found."}), 404

    try:
        base_url = certificate_service.get_public_base_url(request)

        pdf_bytes, filename = certificate_service.generate_certificate_pdf_bytes(
            data=data,
            base_url=base_url,
            template_version=template_version,
        )

        student_name_safe = certificate_service.safe_filename(data.get("student_name", "Student"))
        preview_name = f"Persevex_Certificate_{student_name_safe}_{cert_id}.pdf"

        return Response(
            pdf_bytes,
            mimetype="application/pdf",
            headers={
                "Content-Disposition": f'inline; filename="{preview_name}"',
                "Content-Type": "application/pdf",
                "Cache-Control": "public, max-age=3600",
            },
        )
    except Exception as exc:
        print("CERTIFICATE PREVIEW ERROR:", repr(exc))
        return jsonify({"error": "Certificate document is unavailable for preview."}), 500


@app.get("/verify/<certificate_id>/thumbnail")
@app.get("/verify/<certificate_id>/image")
def verify_image_certificate(certificate_id):
    """
    Publicly accessible endpoint to render and stream a crisp PNG image preview of the certificate.
    Generates image directly in memory on demand from Supabase certificate record or query params.
    NO PERMANENT LOCAL PDF STORAGE REQUIRED.
    """
    cert_id = str(certificate_id).strip()
    record = certificate_service.db_get_certificate_by_id(cert_id, supabase)

    if record:
        if str(record.get("certificate_status") or "").lower() == "revoked":
            return jsonify({"error": "This certificate has been revoked."}), 410
        data = {
            "student_name": record.get("student_name", ""),
            "student_email": record.get("student_email", ""),
            "domain": record.get("internship_domain", ""),
            "start_date": record.get("start_date", ""),
            "end_date": record.get("end_date", ""),
            "issued_date": record.get("issued_date", ""),
            "certificate_id": cert_id,
        }
        template_version = record.get("template_version") or "v1"
    else:
        req_name = request.args.get("name") or request.args.get("student_name")
        req_domain = request.args.get("domain") or request.args.get("internship_domain")
        if req_name or req_domain:
            data = {
                "student_name": req_name or "Student",
                "student_email": request.args.get("email") or request.args.get("student_email") or "",
                "domain": req_domain or "Internship",
                "start_date": request.args.get("start") or request.args.get("start_date") or "",
                "end_date": request.args.get("end") or request.args.get("end_date") or "",
                "issued_date": request.args.get("issued") or request.args.get("issued_date") or "",
                "certificate_id": cert_id,
            }
            template_version = request.args.get("version") or "v1"
        else:
            return jsonify({"error": "Certificate record not found."}), 404

    try:
        base_url = certificate_service.get_public_base_url(request)

        img_bytes = certificate_service.generate_certificate_image_bytes(
            data=data,
            base_url=base_url,
            dpi=180,
            template_version=template_version,
        )

        return Response(
            img_bytes,
            mimetype="image/png",
            headers={
                "Cache-Control": "public, max-age=3600",
                "Content-Type": "image/png",
            },
        )
    except Exception as exc:
        print("CERTIFICATE IMAGE ERROR:", repr(exc))
        return jsonify({"error": "Certificate image could not be rendered."}), 500


# ============================================================
# START SERVER
# ============================================================

if __name__ == "__main__":
    env_info = get_environment_info()
    print("========================================")
    print("Persevex Offer Letter Generator")
    print(f"Environment : {env_info['mode'].upper()}")
    print(f"Supabase Ref: {env_info['supabase_ref']}")
    print("Running at  : http://127.0.0.1:5000")
    print("========================================")

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True,
    )